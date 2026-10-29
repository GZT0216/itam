from __future__ import annotations

import mimetypes
import re
from dataclasses import dataclass
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from fastapi import UploadFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import false, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.asset_ops import InboundStatus
from app.core.assets import ASSET_STATUS_LABELS, AssetStatus, FieldType
from app.core.config import settings
from app.core.exceptions import BizError, ConflictError, NotFoundError
from app.core.permissions import ASSET_FIELD_CATALOG
from app.models import (
    Asset,
    AssetAttachment,
    AssetCategory,
    AssetCategoryField,
    AssetInbound,
    AssetInboundCode,
    AssetModel,
    AuditLog,
    Dept,
    User,
)
from app.models.base import now_cst
from app.repositories.asset_repo import (
    AssetCategoryRepository,
    AssetModelRepository,
    AssetRepository,
)
from app.schemas.asset import (
    AssetAttachmentOut,
    AssetBatchUpdate,
    AssetCreate,
    AssetOut,
    AssetPageOut,
    AssetTimelineItem,
    AssetUpdate,
    CustodianOptionOut,
    DeptOptionOut,
)
from app.services.audit_service import AuditService
from app.services.permission_service import AuthContext
from app.services.scope_service import asset_scope_conditions, is_self_only, scoped_dept_ids
from app.services.txn import commit_or_conflict

_ASSET_FIELD_NAMES = [item["field"] for item in ASSET_FIELD_CATALOG]
_CATALOG_FIELD_SET = set(_ASSET_FIELD_NAMES)
_ASSET_SNAPSHOT_FIELDS = (
    "asset_no",
    "name",
    "category_id",
    "model_id",
    "serial_no",
    "mac_address",
    "purchase_date",
    "supplier",
    "purchase_price",
    "warranty_end",
    "location",
    "dept_id",
    "custodian_id",
    "status",
    "custom_fields",
    "remark",
)


def _asset_snapshot(asset: Asset) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for field_name in _ASSET_SNAPSHOT_FIELDS:
        value = getattr(asset, field_name)
        if field_name == "purchase_price" and value is not None:
            value = float(value)
        result[field_name] = value
    return result


@dataclass
class AssetExport:
    content: bytes
    filename: str
    media_type: str
    row_count: int
    truncated: bool = False


_EXPORT_ROW_LIMIT = 20000

_ATTACHMENT_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}
_ATTACHMENT_FILE_EXTENSIONS = _ATTACHMENT_IMAGE_EXTENSIONS | {
    ".pdf",
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".ppt",
    ".pptx",
    ".txt",
    ".csv",
    ".zip",
    ".rar",
    ".7z",
}
_ATTACHMENT_MAX_SIZE = 10 * 1024 * 1024
_ATTACHMENT_CHUNK_SIZE = 1024 * 1024


class AssetService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.assets = AssetRepository(session)
        self.categories = AssetCategoryRepository(session)
        self.models = AssetModelRepository(session)
        self.audit = AuditService(session)

    async def _scope_conditions(self, context: AuthContext) -> list[Any]:
        return await asset_scope_conditions(self.session, context)

    async def _assert_dept_in_scope(self, dept_id: int | None, context: AuthContext) -> None:
        if is_self_only(context):
            if dept_id is not None and dept_id != context.user.dept_id:
                raise BizError("无权将资产分配到数据范围外的部门")
            return
        allowed = await scoped_dept_ids(self.session, context)
        if allowed is not None and dept_id is not None and dept_id not in allowed:
            raise BizError("无权将资产分配到数据范围外的部门")

    async def _category_ids(
        self, category_id: int | None, include_child: bool
    ) -> list[int] | None:
        if category_id is None:
            return None
        category = await self.categories.get(category_id)
        if category is None:
            raise NotFoundError("资产分类不存在")
        if not include_child:
            return [category.id]
        categories = await self.categories.list_all()
        return [
            item.id
            for item in categories
            if item.id == category.id or item.path.startswith(category.path)
        ]

    def _custom_fields(
        self,
        category: AssetCategory,
        payload: dict[str, Any],
    ) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for field_def in category.field_defs:
            if field_def.status != 1:
                continue
            raw = payload.get(field_def.field_key)
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                if field_def.required:
                    raise BizError(f"自定义字段「{field_def.name}」为必填项")
                continue
            value = self._normalize_field_value(field_def, raw)
            result[field_def.field_key] = value
        return result

    @staticmethod
    def _normalize_field_value(field_def: AssetCategoryField, raw: Any) -> Any:
        label = f"自定义字段「{field_def.name}」"
        if field_def.field_type == FieldType.TEXT:
            text = str(raw).strip()
            if len(text) > 500:
                raise BizError(f"{label}长度不能超过 500 个字符")
            return text
        if field_def.field_type == FieldType.NUMBER:
            if isinstance(raw, bool):
                raise BizError(f"{label}必须是数字")
            try:
                return float(raw)
            except (TypeError, ValueError) as exc:
                raise BizError(f"{label}必须是数字") from exc
        if field_def.field_type == FieldType.DATE:
            try:
                return date.fromisoformat(str(raw)).isoformat()
            except ValueError as exc:
                raise BizError(f"{label}日期格式应为 YYYY-MM-DD") from exc
        if field_def.field_type == FieldType.SELECT:
            text = str(raw).strip()
            if text not in (field_def.options or []):
                raise BizError(f"{label}的取值不在可选范围内")
            return text
        return raw

    async def _resolve_category(self, category_id: int) -> AssetCategory:
        category = await self.categories.get(category_id)
        if category is None:
            raise NotFoundError("资产分类不存在")
        if category.status != 1:
            raise BizError("资产分类已停用，无法录入资产")
        return category

    async def _resolve_model(
        self,
        model_id: int | None,
        category: AssetCategory,
    ) -> AssetModel | None:
        if model_id is None:
            return None
        model = await self.models.get(model_id)
        if model is None:
            raise NotFoundError("资产型号不存在")
        if model.status != 1:
            raise BizError("资产型号已停用，无法录入资产")
        if model.category_id != category.id:
            raise BizError("资产型号与所选分类不一致")
        return model

    async def _resolve_dept(self, dept_id: int | None) -> Dept | None:
        if dept_id is None:
            return None
        dept = await self.session.get(Dept, dept_id)
        if dept is None:
            raise NotFoundError("所属部门不存在")
        if dept.status != 1:
            raise BizError("所属部门已停用")
        return dept

    async def _resolve_custodian(self, custodian_id: int | None) -> User | None:
        if custodian_id is None:
            return None
        stmt = select(User).where(User.id == custodian_id, User.deleted_at.is_(None))
        user = (await self.session.execute(stmt)).scalar_one_or_none()
        if user is None:
            raise NotFoundError("责任人不存在")
        if user.status != 1:
            raise BizError("责任人账号已停用")
        return user

    async def _next_asset_no(self) -> str:
        now = now_cst()
        pattern = settings.ASSET_NO_PATTERN or "{prefix}-{year}{month}-{seq:04d}"
        seq_start = pattern.find("{seq")
        seq_end = pattern.find("}", seq_start)
        if seq_start < 0 or seq_end < 0:
            raise BizError("资产编号规则配置不正确")
        values = {
            "prefix": settings.ASSET_NO_PREFIX,
            "year": now.year,
            "month": f"{now.month:02d}",
        }
        head = pattern[:seq_start].format(**values)
        tail = pattern[seq_end + 1 :]
        max_no = await self.assets.max_asset_no(head)
        sequence = 1
        if max_no:
            middle = max_no[len(head) :]
            if tail and middle.endswith(tail):
                middle = middle[: -len(tail)]
            match = re.match(r"(\d+)", middle)
            if match:
                sequence = int(match.group(1)) + 1
        return pattern.format(**values, seq=sequence)

    def _asset_out(self, asset: Asset, context: AuthContext) -> AssetOut:
        def visible(field_name: str, value: Any) -> Any:
            return value if context.can_view_field("asset", field_name) else None

        masked = [
            name for name in _ASSET_FIELD_NAMES if not context.can_view_field("asset", name)
        ]
        custodian_name = asset.custodian.real_name if asset.custodian else None
        if not context.can_view_field("asset", "custodian_id"):
            custodian_name = None
        return AssetOut(
            id=asset.id,
            asset_no=asset.asset_no,
            name=asset.name if context.can_view_field("asset", "name") else "已隐藏",
            category_id=asset.category_id,
            category_name=asset.category.name if asset.category else None,
            category_type=asset.category.category_type if asset.category else None,
            model_id=asset.model_id,
            model_name=asset.model.name if asset.model else None,
            model_brand=asset.model.brand if asset.model else None,
            serial_no=visible("serial_no", asset.serial_no),
            mac_address=visible("mac_address", asset.mac_address),
            purchase_date=visible("purchase_date", asset.purchase_date),
            supplier=visible("supplier", asset.supplier),
            purchase_price=(
                float(asset.purchase_price) if asset.purchase_price is not None else None
            )
            if context.can_view_field("asset", "purchase_price")
            else None,
            warranty_end=visible("warranty_end", asset.warranty_end),
            location=visible("location", asset.location),
            dept_id=asset.dept_id,
            dept_name=asset.dept.name if asset.dept else None,
            custodian_id=visible("custodian_id", asset.custodian_id),
            custodian_name=custodian_name,
            status=asset.status,
            status_label=ASSET_STATUS_LABELS.get(asset.status, asset.status),
            custom_fields=asset.custom_fields or {},
            remark=visible("remark", asset.remark),
            masked_fields=masked,
            version=asset.version,
            created_at=asset.created_at,
            updated_at=asset.updated_at,
        )

    async def list(
        self,
        context: AuthContext,
        *,
        keyword: str | None,
        category_id: int | None,
        include_child: bool,
        status: str | None,
        dept_id: int | None,
        custodian_id: int | None,
        warranty_days: int | None = None,
        page: int,
        page_size: int,
    ) -> AssetPageOut:
        conditions = await self._list_conditions(
            context,
            keyword=keyword,
            category_id=category_id,
            include_child=include_child,
            status=status,
            dept_id=dept_id,
            custodian_id=custodian_id,
            warranty_days=warranty_days,
        )
        rows, total = await self.assets.list_paginated(
            conditions, page=page, page_size=page_size
        )
        return AssetPageOut(
            items=[self._asset_out(item, context) for item in rows],
            total=total,
            page=page,
            page_size=page_size,
        )

    async def _list_conditions(
        self,
        context: AuthContext,
        *,
        keyword: str | None,
        category_id: int | None,
        include_child: bool,
        status: str | None,
        dept_id: int | None,
        custodian_id: int | None,
        warranty_days: int | None,
    ) -> list[Any]:
        conditions = await self._scope_conditions(context)
        category_ids = await self._category_ids(category_id, include_child)
        if category_ids is not None:
            if not category_ids:
                conditions.append(false())
            else:
                conditions.append(Asset.category_id.in_(category_ids))
        if keyword:
            pattern = f"%{keyword.strip()}%"
            conditions.append(
                or_(
                    Asset.asset_no.like(pattern),
                    Asset.name.like(pattern),
                    Asset.serial_no.like(pattern),
                    Asset.mac_address.like(pattern),
                    Asset.supplier.like(pattern),
                    Asset.location.like(pattern),
                )
            )
        if status:
            if status not in ASSET_STATUS_LABELS:
                raise BizError("资产状态取值不正确")
            conditions.append(Asset.status == status)
        if dept_id is not None:
            conditions.append(Asset.dept_id == dept_id)
        if custodian_id is not None:
            conditions.append(Asset.custodian_id == custodian_id)
        if warranty_days is not None:
            today = now_cst().date()
            conditions.extend(
                [
                    Asset.warranty_end.is_not(None),
                    Asset.warranty_end >= today,
                    Asset.warranty_end <= today + timedelta(days=warranty_days),
                    Asset.status != AssetStatus.SCRAPPED,
                ]
            )
        return conditions

    async def _get_accessible(self, asset_id: int, context: AuthContext) -> Asset:
        asset = await self.assets.get_detail(asset_id)
        if asset is None:
            raise NotFoundError("资产不存在")
        allowed = await scoped_dept_ids(self.session, context)
        if is_self_only(context):
            if asset.custodian_id != context.user.id:
                raise NotFoundError("资产不存在或无权访问")
        elif allowed is not None and asset.dept_id not in allowed:
            raise NotFoundError("资产不存在或无权访问")
        return asset

    async def detail(self, asset_id: int, context: AuthContext) -> AssetOut:
        asset = await self._get_accessible(asset_id, context)
        return self._asset_out(asset, context)

    async def create(self, payload: AssetCreate, context: AuthContext) -> AssetOut:
        if not context.can_edit_field("asset", "name"):
            raise BizError("无权设置资产名称")
        category = await self._resolve_category(payload.category_id)
        model = await self._resolve_model(payload.model_id, category)
        dept = await self._resolve_dept(payload.dept_id)
        custodian = await self._resolve_custodian(payload.custodian_id)
        if dept is None and custodian is not None:
            dept = custodian.dept
            payload.dept_id = custodian.dept_id
        await self._assert_dept_in_scope(dept.id if dept else None, context)
        if is_self_only(context) and (custodian is None or custodian.id != context.user.id):
            raise BizError("当前数据范围仅允许将资产分配给本人")

        asset_no = payload.asset_no
        if asset_no:
            exists = await self.assets.get_by_asset_no(asset_no, include_deleted=True)
            if exists is not None:
                raise ConflictError("资产编号已存在", code=40930)
            reserved = await self.session.scalar(
                select(AssetInboundCode.id)
                .join(
                    AssetInbound,
                    AssetInbound.id == AssetInboundCode.inbound_id,
                )
                .where(
                    AssetInbound.status == InboundStatus.PENDING,
                    func.lower(AssetInboundCode.asset_no) == asset_no.casefold(),
                )
                .limit(1)
            )
            if reserved is not None:
                raise ConflictError(
                    f"资产编号「{asset_no}」已被待确认入库单占用", code=40930
                )
        else:
            asset_no = await self._next_asset_no()

        custom_fields = self._custom_fields(category, payload.custom_fields)
        asset = Asset(
            asset_no=asset_no,
            name=payload.name,
            category_id=category.id,
            category=category,
            model_id=model.id if model else None,
            model=model,
            serial_no=self._editable("serial_no", payload.serial_no, context),
            mac_address=self._editable("mac_address", payload.mac_address, context),
            purchase_date=self._editable("purchase_date", payload.purchase_date, context),
            supplier=self._editable("supplier", payload.supplier, context),
            purchase_price=self._editable("purchase_price", payload.purchase_price, context),
            warranty_end=self._editable("warranty_end", payload.warranty_end, context),
            location=self._editable("location", payload.location, context),
            dept_id=dept.id if dept else None,
            dept=dept,
            custodian_id=custodian.id if custodian else None,
            custodian=custodian,
            status=AssetStatus.PENDING_STOCK,
            custom_fields=custom_fields,
            remark=self._editable("remark", payload.remark, context),
        )
        self.session.add(asset)
        await self.session.flush()
        await self.audit.record(
            action="asset.ledger.create",
            module="资产台账",
            summary=f"新增资产：{asset.name}（{asset.asset_no}）",
            actor=context.user,
            resource_type="asset",
            resource_id=asset.id,
            after=_asset_snapshot(asset),
        )
        await commit_or_conflict(self.session)
        return self._asset_out(asset, context)

    def _editable(self, field_name: str, value: Any, context: AuthContext) -> Any:
        if field_name not in _CATALOG_FIELD_SET:
            return value
        return value if context.can_edit_field("asset", field_name) else None

    async def update(
        self,
        asset_id: int,
        payload: AssetUpdate,
        context: AuthContext,
    ) -> AssetOut:
        asset = await self._get_accessible(asset_id, context)
        if asset.version != payload.version:
            raise ConflictError("资产信息已被他人修改，请刷新后重试")
        category = await self._resolve_category(payload.category_id)
        model = await self._resolve_model(payload.model_id, category)
        dept = await self._resolve_dept(payload.dept_id)
        custodian = await self._resolve_custodian(payload.custodian_id)
        if dept is None and custodian is not None:
            dept = custodian.dept
            payload.dept_id = custodian.dept_id
        await self._assert_dept_in_scope(dept.id if dept else None, context)
        if is_self_only(context) and asset.custodian_id != context.user.id:
            raise BizError("无权修改数据范围外的资产")

        before = _asset_snapshot(asset)
        asset.name = self._keep_if_locked("name", payload.name, asset.name, context)
        asset.category_id = category.id
        asset.category = category
        asset.model_id = model.id if model else None
        asset.model = model
        asset.serial_no = self._keep_if_locked(
            "serial_no", payload.serial_no, asset.serial_no, context
        )
        asset.mac_address = self._keep_if_locked(
            "mac_address", payload.mac_address, asset.mac_address, context
        )
        asset.purchase_date = self._keep_if_locked(
            "purchase_date", payload.purchase_date, asset.purchase_date, context
        )
        asset.supplier = self._keep_if_locked(
            "supplier", payload.supplier, asset.supplier, context
        )
        asset.purchase_price = self._keep_if_locked(
            "purchase_price", payload.purchase_price, asset.purchase_price, context
        )
        asset.warranty_end = self._keep_if_locked(
            "warranty_end", payload.warranty_end, asset.warranty_end, context
        )
        asset.location = self._keep_if_locked(
            "location", payload.location, asset.location, context
        )
        asset.dept_id = dept.id if dept else None
        asset.dept = dept
        asset.custodian_id = self._keep_if_locked(
            "custodian_id", custodian.id if custodian else None, asset.custodian_id, context
        )
        asset.custodian = custodian if asset.custodian_id == (custodian.id if custodian else None) else asset.custodian
        asset.remark = self._keep_if_locked("remark", payload.remark, asset.remark, context)
        asset.custom_fields = self._custom_fields(category, payload.custom_fields)
        await self.audit.record(
            action="asset.ledger.update",
            module="资产台账",
            summary=f"编辑资产：{asset.name}（{asset.asset_no}）",
            actor=context.user,
            resource_type="asset",
            resource_id=asset.id,
            before=before,
            after=_asset_snapshot(asset),
        )
        await commit_or_conflict(self.session)
        return self._asset_out(asset, context)

    def _keep_if_locked(
        self,
        field_name: str,
        new_value: Any,
        old_value: Any,
        context: AuthContext,
    ) -> Any:
        if field_name not in _CATALOG_FIELD_SET:
            return new_value
        return new_value if context.can_edit_field("asset", field_name) else old_value

    async def delete(self, asset_id: int, context: AuthContext) -> None:
        asset = await self._get_accessible(asset_id, context)
        if asset.status not in (AssetStatus.PENDING_STOCK, AssetStatus.IN_STOCK):
            raise BizError("仅待入库或在库资产可直接删除，其他状态请先走业务流程")
        before = _asset_snapshot(asset)
        asset.deleted_at = now_cst()
        await self.audit.record(
            action="asset.ledger.delete",
            module="资产台账",
            summary=f"删除资产：{asset.name}（{asset.asset_no}）",
            actor=context.user,
            resource_type="asset",
            resource_id=asset.id,
            before=before,
        )
        await commit_or_conflict(self.session)

    async def batch_update(self, payload: AssetBatchUpdate, context: AuthContext) -> int:
        if (
            payload.custodian_id is not None or payload.clear_custodian
        ) and not context.can_edit_field("asset", "custodian_id"):
            raise BizError("无权批量修改资产责任人")
        if (
            payload.location is not None or payload.clear_location
        ) and not context.can_edit_field("asset", "location"):
            raise BizError("无权批量修改存放位置")
        dept = await self._resolve_dept(payload.dept_id)
        custodian = await self._resolve_custodian(payload.custodian_id)
        await self._assert_dept_in_scope(dept.id if dept else None, context)
        if is_self_only(context):
            if payload.custodian_id is not None and payload.custodian_id != context.user.id:
                raise BizError("当前数据范围仅允许将资产分配给本人")
            if payload.clear_custodian:
                raise BizError("当前数据范围不允许清除资产责任人")

        conditions = await self._scope_conditions(context)
        conditions.extend(
            [Asset.id.in_(payload.asset_ids), Asset.deleted_at.is_(None)]
        )
        rows = list(
            (await self.session.execute(select(Asset).where(*conditions)))
            .scalars()
            .all()
        )
        if len(rows) != len(payload.asset_ids):
            raise NotFoundError("部分资产不存在或不在数据范围内")

        target_dept = dept
        if target_dept is None and custodian is not None and not payload.clear_dept:
            target_dept = custodian.dept

        for asset in rows:
            before = _asset_snapshot(asset)
            if target_dept is not None:
                asset.dept_id = target_dept.id
                asset.dept = target_dept
            elif payload.clear_dept or custodian is not None:
                asset.dept_id = None
                asset.dept = None
            if custodian is not None:
                asset.custodian_id = custodian.id
                asset.custodian = custodian
            elif payload.clear_custodian:
                asset.custodian_id = None
                asset.custodian = None
            if payload.location is not None:
                asset.location = payload.location
            elif payload.clear_location:
                asset.location = None
            await self.audit.record(
                action="asset.ledger.update",
                module="资产台账",
                summary=f"批量编辑资产：{asset.name}（{asset.asset_no}）",
                actor=context.user,
                resource_type="asset",
                resource_id=asset.id,
                before=before,
                after=_asset_snapshot(asset),
            )
        await commit_or_conflict(self.session)
        return len(rows)

    async def export(
        self,
        context: AuthContext,
        *,
        keyword: str | None,
        category_id: int | None,
        include_child: bool,
        status: str | None,
        dept_id: int | None,
        custodian_id: int | None,
        warranty_days: int | None,
    ) -> AssetExport:
        conditions = await self._list_conditions(
            context,
            keyword=keyword,
            category_id=category_id,
            include_child=include_child,
            status=status,
            dept_id=dept_id,
            custodian_id=custodian_id,
            warranty_days=warranty_days,
        )
        stmt = (
            select(Asset)
            .where(*conditions)
            .order_by(Asset.id.desc())
            .limit(_EXPORT_ROW_LIMIT + 1)
        )
        rows = list((await self.session.execute(stmt)).scalars().all())
        truncated = len(rows) > _EXPORT_ROW_LIMIT
        if truncated:
            rows = rows[:_EXPORT_ROW_LIMIT]

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "资产台账"
        sheet.append(
            [
                "资产编号",
                "资产名称",
                "资产分类",
                "型号",
                "序列号",
                "MAC 地址",
                "资产状态",
                "所属部门",
                "责任人",
                "存放位置",
                "购置日期",
                "采购价格",
                "保修到期",
                "供应商",
                "备注",
            ]
        )
        for row in rows:
            item = self._asset_out(row, context)
            sheet.append(
                [
                    item.asset_no,
                    item.name,
                    item.category_name or "",
                    item.model_name or "",
                    item.serial_no or "",
                    item.mac_address or "",
                    item.status_label,
                    item.dept_name or "",
                    item.custodian_name or "",
                    item.location or "",
                    item.purchase_date.isoformat() if item.purchase_date else "",
                    item.purchase_price if item.purchase_price is not None else "",
                    item.warranty_end.isoformat() if item.warranty_end else "",
                    item.supplier or "",
                    item.remark or "",
                ]
            )

        header_fill = PatternFill("solid", fgColor="DCE6F1")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        widths = [22, 26, 18, 22, 20, 20, 12, 18, 14, 20, 14, 14, 14, 22, 30]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet["L"][1:]:
            cell.number_format = "#,##0.00"

        buffer = BytesIO()
        workbook.save(buffer)
        filename = f"资产台账_{now_cst().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return AssetExport(
            content=buffer.getvalue(),
            filename=filename,
            media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            row_count=len(rows),
            truncated=truncated,
        )

    async def timeline(self, asset_id: int, context: AuthContext) -> list[AssetTimelineItem]:
        await self._get_accessible(asset_id, context)
        stmt = (
            select(AuditLog)
            .where(
                AuditLog.resource_type == "asset",
                AuditLog.resource_id == str(asset_id),
            )
            .order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
            .limit(100)
        )
        logs = list((await self.session.execute(stmt)).scalars().all())
        return [
            AssetTimelineItem(
                id=item.id,
                action=item.action,
                summary=item.summary,
                username=item.username,
                status=item.status,
                created_at=item.created_at,
            )
            for item in logs
        ]

    @staticmethod
    def _attachment_out(item: AssetAttachment) -> AssetAttachmentOut:
        return AssetAttachmentOut(
            id=item.id,
            asset_id=item.asset_id,
            file_name=item.file_name,
            content_type=item.content_type,
            file_size=item.file_size,
            kind=item.kind,
            uploaded_by_name=item.uploaded_by_name,
            created_at=item.created_at,
        )

    @staticmethod
    def _attachment_path(attachment: AssetAttachment) -> Path:
        return (
            Path(settings.UPLOAD_DIR)
            / "assets"
            / str(attachment.asset_id)
            / attachment.stored_name
        )

    async def list_attachments(
        self, asset_id: int, context: AuthContext
    ) -> list[AssetAttachmentOut]:
        await self._get_accessible(asset_id, context)
        stmt = (
            select(AssetAttachment)
            .where(
                AssetAttachment.asset_id == asset_id,
                AssetAttachment.deleted_at.is_(None),
            )
            .order_by(AssetAttachment.id.desc())
        )
        rows = list((await self.session.execute(stmt)).scalars().all())
        return [self._attachment_out(item) for item in rows]

    async def upload_attachment(
        self, asset_id: int, upload: UploadFile, context: AuthContext
    ) -> AssetAttachmentOut:
        asset = await self._get_accessible(asset_id, context)
        original_name = Path(upload.filename or "").name.strip()
        if not original_name:
            raise BizError("请选择要上传的文件")
        if len(original_name) > 255:
            raise BizError("文件名不能超过 255 个字符")
        extension = Path(original_name).suffix.lower()
        if extension not in _ATTACHMENT_FILE_EXTENSIONS:
            allowed = "、".join(sorted(item.lstrip(".") for item in _ATTACHMENT_FILE_EXTENSIONS))
            raise BizError(f"不支持的文件类型，可上传：{allowed}")

        chunks: list[bytes] = []
        size = 0
        while True:
            chunk = await upload.read(_ATTACHMENT_CHUNK_SIZE)
            if not chunk:
                break
            size += len(chunk)
            if size > _ATTACHMENT_MAX_SIZE:
                raise BizError("单个附件不能超过 10MB")
            chunks.append(chunk)
        if size == 0:
            raise BizError("文件内容为空")

        content_type = (
            upload.content_type
            or mimetypes.guess_type(original_name)[0]
            or "application/octet-stream"
        )
        stored_name = f"{uuid4().hex}{extension}"
        directory = Path(settings.UPLOAD_DIR) / "assets" / str(asset.id)
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / stored_name
        path.write_bytes(b"".join(chunks))

        attachment = AssetAttachment(
            asset_id=asset.id,
            file_name=original_name,
            stored_name=stored_name,
            content_type=content_type,
            file_size=size,
            kind="IMAGE" if extension in _ATTACHMENT_IMAGE_EXTENSIONS else "FILE",
            uploaded_by_id=context.user.id,
            uploaded_by_name=context.user.real_name,
        )
        self.session.add(attachment)
        await self.session.flush()
        await self.audit.record(
            action="asset.attachment.upload",
            module="资产台账",
            summary=f"上传资产附件：{original_name}（{asset.asset_no}）",
            actor=context.user,
            resource_type="asset",
            resource_id=asset.id,
            after={
                "file_name": original_name,
                "file_size": size,
                "kind": attachment.kind,
            },
        )
        try:
            await commit_or_conflict(self.session)
        except Exception:
            path.unlink(missing_ok=True)
            raise
        return self._attachment_out(attachment)

    async def get_attachment_file(
        self, asset_id: int, attachment_id: int, context: AuthContext
    ) -> tuple[Path, AssetAttachment]:
        await self._get_accessible(asset_id, context)
        attachment = await self.session.scalar(
            select(AssetAttachment).where(
                AssetAttachment.id == attachment_id,
                AssetAttachment.asset_id == asset_id,
                AssetAttachment.deleted_at.is_(None),
            )
        )
        if attachment is None:
            raise NotFoundError("附件不存在")
        path = self._attachment_path(attachment)
        if not path.exists():
            raise NotFoundError("附件文件已丢失，请联系管理员")
        return path, attachment

    async def delete_attachment(
        self, asset_id: int, attachment_id: int, context: AuthContext
    ) -> None:
        await self._get_accessible(asset_id, context)
        attachment = await self.session.scalar(
            select(AssetAttachment).where(
                AssetAttachment.id == attachment_id,
                AssetAttachment.asset_id == asset_id,
                AssetAttachment.deleted_at.is_(None),
            )
        )
        if attachment is None:
            raise NotFoundError("附件不存在")
        path = self._attachment_path(attachment)
        attachment.deleted_at = now_cst()
        await self.audit.record(
            action="asset.attachment.delete",
            module="资产台账",
            summary=f"删除资产附件：{attachment.file_name}",
            actor=context.user,
            resource_type="asset",
            resource_id=asset_id,
            before={
                "file_name": attachment.file_name,
                "file_size": attachment.file_size,
            },
        )
        await commit_or_conflict(self.session)
        path.unlink(missing_ok=True)

    async def custodian_options(
        self, context: AuthContext, *, keyword: str | None
    ) -> list[CustodianOptionOut]:
        stmt = select(User).where(User.deleted_at.is_(None), User.status == 1)
        if keyword:
            pattern = f"%{keyword.strip()}%"
            stmt = stmt.where(
                or_(
                    User.username.like(pattern),
                    User.real_name.like(pattern),
                    User.employee_no.like(pattern),
                )
            )
        allowed = await scoped_dept_ids(self.session, context)
        if allowed is not None:
            if not allowed:
                return []
            stmt = stmt.where(User.dept_id.in_(allowed))
        stmt = stmt.order_by(User.real_name, User.id).limit(50)
        users = list((await self.session.execute(stmt)).scalars().all())
        return [
            CustodianOptionOut(
                id=user.id,
                username=user.username,
                real_name=user.real_name,
                dept_id=user.dept_id,
                dept_name=user.dept.name if user.dept else None,
            )
            for user in users
        ]

    async def dept_options(self, context: AuthContext) -> list[DeptOptionOut]:
        stmt = select(Dept).where(Dept.status == 1).order_by(Dept.path, Dept.sort, Dept.id)
        allowed = await scoped_dept_ids(self.session, context)
        if is_self_only(context):
            allowed = {context.user.dept_id} if context.user.dept_id else set()
        if allowed is not None:
            if not allowed:
                return []
            stmt = stmt.where(Dept.id.in_(allowed))
        depts = list((await self.session.execute(stmt)).scalars().all())
        return [
            DeptOptionOut(
                id=dept.id,
                parent_id=dept.parent_id,
                name=dept.name,
                path=dept.path,
                level=dept.level,
            )
            for dept in depts
        ]
