from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import BizError
from app.models import AssetCategory, AssetModel, Dept, Role, User
from app.repositories.user_repo import UserRepository
from app.schemas.asset import AssetCreate, CategoryCreate
from app.schemas.org import UserCreate
from app.services.asset_catalog_service import CategoryService
from app.services.asset_service import AssetService
from app.services.permission_service import AuthContext
from app.services.user_admin_service import UserAdminService

MAX_IMPORT_ROWS = 1000
DEFAULT_USER_PASSWORD = "Itam@123456"

CATEGORY_HEADERS = [
    "分类编码",
    "分类名称",
    "上级分类编码",
    "分类类型",
    "排序",
    "安全库存",
    "状态",
    "备注",
]
USER_HEADERS = [
    "用户名",
    "姓名",
    "部门编码",
    "角色编码",
    "初始密码",
    "工号",
    "邮箱",
    "手机号",
    "职位",
    "状态",
    "备注",
]
ASSET_HEADERS = [
    "资产名称",
    "分类编码",
    "资产编号",
    "型号编码",
    "序列号",
    "MAC 地址",
    "购置日期",
    "供应商",
    "采购价格",
    "保修到期",
    "存放位置",
    "部门编码",
    "责任人用户名",
    "备注",
]

_CATEGORY_TYPE_MAP = {
    "硬件": "HARDWARE",
    "软件": "SOFTWARE",
    "配件": "ACCESSORY",
    "耗材": "CONSUMABLE",
    "HARDWARE": "HARDWARE",
    "SOFTWARE": "SOFTWARE",
    "ACCESSORY": "ACCESSORY",
    "CONSUMABLE": "CONSUMABLE",
}
_IMAGE_NOTES = [
    "请勿修改表头名称，否则无法识别数据列。",
    "第 2 行为示例数据，导入前请删除或替换为真实数据。",
    "单次最多导入 1000 行，单文件不超过 10MB。",
]


@dataclass
class ImportResult:
    total: int
    success: int
    failed: int
    errors: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class ImportTemplate:
    content: bytes
    filename: str
    media_type: str = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def _date_text(value: Any, label: str) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    try:
        return date.fromisoformat(text).isoformat()
    except ValueError as exc:
        raise ValueError(f"{label}格式应为 YYYY-MM-DD") from exc


def _int_value(value: Any, default: int = 0) -> int:
    text = _text(value)
    if text is None:
        return default
    try:
        number = float(text)
    except ValueError as exc:
        raise ValueError(f"「{text}」不是有效的整数") from exc
    if number != int(number):
        raise ValueError(f"「{text}」不是有效的整数")
    return int(number)


def _price_value(value: Any) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError("采购价格必须为数字") from exc
    if number < 0:
        raise ValueError("采购价格不能为负数")
    return number


def _status_value(value: Any, default: int = 1) -> int:
    text = _text(value)
    if text is None:
        return default
    if text in {"启用", "正常", "是", "1", "true", "TRUE"}:
        return 1
    if text in {"停用", "禁用", "否", "0", "false", "FALSE"}:
        return 0
    raise ValueError("状态仅支持「启用」或「停用」")


def _category_type(value: Any) -> str | None:
    text = _text(value)
    if text is None:
        return None
    result = _CATEGORY_TYPE_MAP.get(text.upper() if text.isascii() else text)
    if result is None:
        raise ValueError("分类类型仅支持：硬件 / 软件 / 配件 / 耗材")
    return result


def _error_message(exc: Exception) -> str:
    if isinstance(exc, ValidationError):
        first = exc.errors()[0]
        message = str(first.get("msg", "参数错误"))
        location = ".".join(
            str(part) for part in first.get("loc", ()) if part not in ("body",)
        )
        return f"{location}: {message}" if location else message
    return str(getattr(exc, "message", exc) or exc)


def _style_header(sheet) -> None:
    fill = PatternFill("solid", fgColor="DCE6F1")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


class ImportService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self._actor_id: int | None = None

    async def _rollback_and_restore(self, context: AuthContext) -> None:
        # 回滚会过期当前登录用户，后续审计写入需要重新加载
        actor_id = self._actor_id or context.user.id
        await self.session.rollback()
        user = await UserRepository(self.session).get_with_relations(actor_id)
        if user is not None:
            context.user = user

    @staticmethod
    def _read_rows(content: bytes, headers: list[str]) -> list[tuple[int, dict[str, Any]]]:
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise BizError("无法解析 Excel 文件，请使用 .xlsx 模板") from exc
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration as exc:
                raise BizError("Excel 文件为空") from exc
            normalized = [
                str(cell).strip() if cell is not None else "" for cell in header_row
            ]
            missing = [name for name in headers if name not in normalized]
            if missing:
                raise BizError(
                    f"表头缺少字段：{'、'.join(missing)}，请下载最新模板"
                )
            index_map = {name: normalized.index(name) for name in headers}
            data: list[tuple[int, dict[str, Any]]] = []
            for row_number, values in enumerate(rows, start=2):
                if all(value is None or str(value).strip() == "" for value in values):
                    continue
                data.append(
                    (
                        row_number,
                        {
                            name: values[index] if index < len(values) else None
                            for name, index in index_map.items()
                        },
                    )
                )
                if len(data) > MAX_IMPORT_ROWS:
                    raise BizError(
                        f"单次最多导入 {MAX_IMPORT_ROWS} 行，请拆分后重试"
                    )
            if not data:
                raise BizError("Excel 中没有可导入的数据行")
            return data
        finally:
            workbook.close()

    async def _dept_map(self) -> dict[str, int]:
        rows = (await self.session.execute(select(Dept.id, Dept.code))).all()
        return {code.upper(): dept_id for dept_id, code in rows}

    async def _category_map(self) -> dict[str, int]:
        rows = (
            await self.session.execute(select(AssetCategory.id, AssetCategory.code))
        ).all()
        return {code.upper(): category_id for category_id, code in rows}

    async def _model_map(self) -> dict[str, int]:
        rows = (
            await self.session.execute(select(AssetModel.id, AssetModel.code))
        ).all()
        return {code.upper(): model_id for model_id, code in rows}

    async def _user_map(self) -> dict[str, int]:
        rows = (
            await self.session.execute(
                select(User.id, User.username).where(User.deleted_at.is_(None))
            )
        ).all()
        return {username.lower(): user_id for user_id, username in rows}

    async def _role_map(self) -> dict[str, int]:
        rows = (await self.session.execute(select(Role.id, Role.code))).all()
        return {code.upper(): role_id for role_id, code in rows}

    @staticmethod
    def _reject(row: int, message: str, result: ImportResult) -> None:
        result.failed += 1
        result.errors.append({"row": row, "message": message})

    async def import_categories(
        self, content: bytes, context: AuthContext
    ) -> ImportResult:
        self._actor_id = context.user.id
        rows = self._read_rows(content, CATEGORY_HEADERS)
        result = ImportResult(total=len(rows), success=0, failed=0)
        code_map = await self._category_map()
        service = CategoryService(self.session)

        pending = list(rows)
        while pending:
            progressed = False
            next_pending: list[tuple[int, dict[str, Any]]] = []
            for row_number, row in pending:
                code = (_text(row["分类编码"]) or "").upper()
                parent_code = (_text(row["上级分类编码"]) or "").upper()
                if parent_code and parent_code not in code_map:
                    next_pending.append((row_number, row))
                    continue
                try:
                    created = await service.create(
                        CategoryCreate(
                            parent_id=code_map[parent_code] if parent_code else None,
                            name=_text(row["分类名称"]) or "",
                            code=code,
                            category_type=_category_type(row["分类类型"]),
                            sort=_int_value(row["排序"], 0),
                            safety_stock=_int_value(row["安全库存"], 0),
                            status=_status_value(row["状态"], 1),
                            remark=_text(row["备注"]),
                        ),
                        context,
                    )
                    code_map[created.code.upper()] = created.id
                    result.success += 1
                    progressed = True
                except Exception as exc:
                    await self._rollback_and_restore(context)
                    self._reject(row_number, _error_message(exc), result)
                    progressed = True
            if not progressed:
                for row_number, row in next_pending:
                    parent_code = (_text(row["上级分类编码"]) or "").upper()
                    self._reject(
                        row_number,
                        f"上级分类编码「{parent_code}」不存在",
                        result,
                    )
                break
            pending = next_pending
        return result

    async def import_users(self, content: bytes, context: AuthContext) -> ImportResult:
        self._actor_id = context.user.id
        rows = self._read_rows(content, USER_HEADERS)
        result = ImportResult(total=len(rows), success=0, failed=0)
        dept_map = await self._dept_map()
        role_map = await self._role_map()
        service = UserAdminService(self.session)

        for row_number, row in rows:
            try:
                dept_code = (_text(row["部门编码"]) or "").upper()
                dept_id = dept_map.get(dept_code)
                if dept_id is None:
                    raise ValueError(f"部门编码「{dept_code}」不存在")
                role_codes = [
                    item
                    for item in re.split(
                        r"[,，;；、\s]+", _text(row["角色编码"]) or ""
                    )
                    if item
                ]
                if not role_codes:
                    role_codes = ["EMPLOYEE"]
                role_ids: list[int] = []
                for code in role_codes:
                    role_id = role_map.get(code.upper())
                    if role_id is None:
                        raise ValueError(f"角色编码「{code}」不存在")
                    role_ids.append(role_id)
                password = _text(row["初始密码"]) or DEFAULT_USER_PASSWORD
                await service.create_user(
                    UserCreate(
                        username=_text(row["用户名"]) or "",
                        password=password,
                        real_name=_text(row["姓名"]) or "",
                        dept_id=dept_id,
                        position=_text(row["职位"]),
                        employee_no=_text(row["工号"]),
                        email=_text(row["邮箱"]),
                        phone=_text(row["手机号"]),
                        status=_status_value(row["状态"], 1),
                        role_ids=role_ids,
                        remark=_text(row["备注"]),
                    ),
                    context,
                )
                result.success += 1
            except Exception as exc:
                await self._rollback_and_restore(context)
                self._reject(row_number, _error_message(exc), result)
        return result

    async def import_assets(self, content: bytes, context: AuthContext) -> ImportResult:
        self._actor_id = context.user.id
        rows = self._read_rows(content, ASSET_HEADERS)
        result = ImportResult(total=len(rows), success=0, failed=0)
        category_map = await self._category_map()
        model_map = await self._model_map()
        dept_map = await self._dept_map()
        user_map = await self._user_map()
        service = AssetService(self.session)

        for row_number, row in rows:
            try:
                category_code = (_text(row["分类编码"]) or "").upper()
                category_id = category_map.get(category_code)
                if category_id is None:
                    raise ValueError(f"分类编码「{category_code}」不存在")
                model_code = (_text(row["型号编码"]) or "").upper()
                model_id = model_map.get(model_code) if model_code else None
                if model_code and model_id is None:
                    raise ValueError(f"型号编码「{model_code}」不存在")
                dept_code = (_text(row["部门编码"]) or "").upper()
                dept_id = dept_map.get(dept_code) if dept_code else None
                if dept_code and dept_id is None:
                    raise ValueError(f"部门编码「{dept_code}」不存在")
                username = (_text(row["责任人用户名"]) or "").lower()
                custodian_id = user_map.get(username) if username else None
                if username and custodian_id is None:
                    raise ValueError(f"责任人用户名「{username}」不存在")
                await service.create(
                    AssetCreate(
                        asset_no=_text(row["资产编号"]),
                        name=_text(row["资产名称"]) or "",
                        category_id=category_id,
                        model_id=model_id,
                        serial_no=_text(row["序列号"]),
                        mac_address=_text(row["MAC 地址"]),
                        purchase_date=_date_text(row["购置日期"], "购置日期"),
                        supplier=_text(row["供应商"]),
                        purchase_price=_price_value(row["采购价格"]),
                        warranty_end=_date_text(row["保修到期"], "保修到期"),
                        location=_text(row["存放位置"]),
                        dept_id=dept_id,
                        custodian_id=custodian_id,
                        remark=_text(row["备注"]),
                        custom_fields={},
                    ),
                    context,
                )
                result.success += 1
            except Exception as exc:
                await self._rollback_and_restore(context)
                self._reject(row_number, _error_message(exc), result)
        return result

    @staticmethod
    def template_categories() -> ImportTemplate:
        return ImportService._build_template(
            "资产分类导入模板",
            CATEGORY_HEADERS,
            ["HW-NOTEBOOK", "笔记本电脑", "", "硬件", 10, 0, "启用", "示例数据，导入前请删除"],
            [
                *_IMAGE_NOTES,
                "分类编码仅支持大写字母、数字、下划线或短横线，且不能重复。",
                "上级分类编码留空表示创建根分类；根分类必须填写分类类型。",
                "分类类型可选：硬件 / 软件 / 配件 / 耗材。",
                "本模板不支持导入自定义字段，可在分类创建后单独维护。",
            ],
        )

    @staticmethod
    def template_users() -> ImportTemplate:
        return ImportService._build_template(
            "用户导入模板",
            USER_HEADERS,
            [
                "zhangsan",
                "张三",
                "IT",
                "EMPLOYEE",
                DEFAULT_USER_PASSWORD,
                "E1001",
                "zhangsan@example.com",
                "13800000000",
                "工程师",
                "启用",
                "示例数据，导入前请删除",
            ],
            [
                *_IMAGE_NOTES,
                "用户名需为 3-50 位小写字母、数字、点、下划线或短横线，且不能重复。",
                "部门编码、角色编码必须已存在于系统中；角色编码可用逗号分隔多个。",
                f"初始密码留空时使用默认密码 {DEFAULT_USER_PASSWORD}，首次登录会强制修改。",
                "状态可选：启用 / 停用。",
            ],
        )

    @staticmethod
    def template_assets() -> ImportTemplate:
        return ImportService._build_template(
            "资产导入模板",
            ASSET_HEADERS,
            [
                "示例笔记本",
                "HW-LAPTOP",
                "",
                "",
                "SN-DEMO-001",
                "AA:BB:CC:DD:EE:FF",
                "2026-01-15",
                "示例供应商",
                8800,
                "2029-01-14",
                "A 座 301",
                "IT",
                "admin",
                "示例数据，导入前请删除",
            ],
            [
                *_IMAGE_NOTES,
                "分类编码必填且需已存在；型号编码、资产编号选填。",
                "资产编号留空时按系统规则自动生成；型号需属于所选分类。",
                "购置日期、保修到期格式为 YYYY-MM-DD。",
                "部门编码、责任人用户名需已存在；导入资产初始状态为「待入库」。",
                "本模板不支持导入分类自定义字段，缺失必填扩展字段的行会提示错误。",
            ],
        )

    @staticmethod
    def _build_template(
        title: str,
        headers: list[str],
        example: list[Any],
        notes: list[str],
    ) -> ImportTemplate:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "导入数据"
        sheet.append(headers)
        sheet.append(example)
        _style_header(sheet)
        widths = [18, 20, 16, 16, 14, 16, 14, 14, 14, 14, 20, 16, 16, 30]
        for index in range(1, len(headers) + 1):
            width = widths[index - 1] if index <= len(widths) else 18
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"

        guide = workbook.create_sheet("填写说明")
        guide.append(["填写说明"])
        guide["A1"].font = Font(bold=True, size=13)
        for note in notes:
            guide.append([note])
        guide.column_dimensions["A"].width = 90
        for cell in guide["A"]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        buffer = BytesIO()
        workbook.save(buffer)
        return ImportTemplate(
            content=buffer.getvalue(),
            filename=f"{title}.xlsx",
        )
