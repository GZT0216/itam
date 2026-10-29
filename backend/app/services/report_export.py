from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Iterable
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.schemas.report import (
    AssetReportOverviewOut,
    ComplianceAssetRow,
    ComplianceIssueSummary,
    IdleAssetRow,
    ReportAssetRow,
)

_EXCEL_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_EXCEL_TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
_EXCEL_META_FONT = Font(size=9, color="666666")
_EXCEL_HEADER_FONT = Font(size=10, bold=True, color="FFFFFF")
_EXCEL_BODY_FONT = Font(size=9)
_EXCEL_THIN = Side(style="thin", color="D9D9D9")
_EXCEL_BORDER = Border(left=_EXCEL_THIN, right=_EXCEL_THIN, top=_EXCEL_THIN, bottom=_EXCEL_THIN)
_EXCEL_CENTER = Alignment(horizontal="center", vertical="center")
_EXCEL_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)

_PDF_FONT = "STSong-Light"
pdfmetrics.registerFont(UnicodeCIDFont(_PDF_FONT))

_PDF_TITLE = ParagraphStyle(
    "ItamTitle",
    fontName=_PDF_FONT,
    fontSize=18,
    leading=24,
    alignment=TA_CENTER,
    textColor=colors.HexColor("#1F4E78"),
    spaceAfter=8,
)
_PDF_META = ParagraphStyle(
    "ItamMeta",
    fontName=_PDF_FONT,
    fontSize=8.5,
    leading=13,
    alignment=TA_LEFT,
    textColor=colors.HexColor("#555555"),
)
_PDF_HEADING = ParagraphStyle(
    "ItamHeading",
    fontName=_PDF_FONT,
    fontSize=11,
    leading=16,
    textColor=colors.HexColor("#1F4E78"),
    spaceBefore=10,
    spaceAfter=5,
)
_PDF_CELL = ParagraphStyle(
    "ItamCell",
    fontName=_PDF_FONT,
    fontSize=7.5,
    leading=10,
    alignment=TA_LEFT,
)
_PDF_CELL_CENTER = ParagraphStyle(
    "ItamCellCenter",
    parent=_PDF_CELL,
    alignment=TA_CENTER,
)
_PDF_HEADER_CELL = ParagraphStyle(
    "ItamHeaderCell",
    parent=_PDF_CELL_CENTER,
    textColor=colors.white,
)


def _safe_excel(value: Any) -> Any:
    if isinstance(value, str):
        if value[:1] in ("=", "+", "-", "@"):
            return "'" + value
        return value
    if isinstance(value, Decimal):
        return float(value)
    return value


def _fmt_date(value: date | datetime | None) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    return value.isoformat()


def _fmt_money(value: float | Decimal | None) -> str:
    if value is None:
        return ""
    return f"{float(value):,.2f}"


def _write_intro(
    ws,
    title: str,
    filter_text: str,
    data_scope: str,
    generated_at: datetime,
    column_count: int,
) -> int:
    last_column = get_column_letter(column_count)
    ws.merge_cells(f"A1:{last_column}1")
    title_cell = ws["A1"]
    title_cell.value = f"ITAM {title}"
    title_cell.font = Font(size=15, bold=True, color="FFFFFF")
    title_cell.fill = _EXCEL_TITLE_FILL
    title_cell.alignment = _EXCEL_CENTER
    ws.row_dimensions[1].height = 28

    ws.merge_cells(f"A2:{last_column}2")
    ws["A2"] = f"筛选条件：{filter_text}"
    ws["A2"].font = _EXCEL_META_FONT
    ws["A2"].alignment = _EXCEL_LEFT

    ws.merge_cells(f"A3:{last_column}3")
    generated = generated_at.strftime("%Y-%m-%d %H:%M:%S")
    ws["A3"] = f"数据范围：{data_scope}    生成时间：{generated}"
    ws["A3"].font = _EXCEL_META_FONT
    ws["A3"].alignment = _EXCEL_LEFT
    return 5


def _write_table(
    ws,
    headers: list[str],
    rows: Iterable[Iterable[Any]],
    *,
    start_row: int,
    widths: list[int] | None = None,
    money_columns: set[int] | None = None,
    date_columns: set[int] | None = None,
    integer_columns: set[int] | None = None,
) -> int:
    money_columns = money_columns or set()
    date_columns = date_columns or set()
    integer_columns = integer_columns or set()

    for index, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=index, value=header)
        cell.font = _EXCEL_HEADER_FONT
        cell.fill = _EXCEL_HEADER_FILL
        cell.alignment = _EXCEL_CENTER
        cell.border = _EXCEL_BORDER
        if widths:
            ws.column_dimensions[get_column_letter(index)].width = widths[index - 1]

    row_index = start_row
    for values in rows:
        row_index += 1
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=column, value=_safe_excel(value))
            cell.font = _EXCEL_BODY_FONT
            cell.border = _EXCEL_BORDER
            cell.alignment = _EXCEL_CENTER if column in money_columns or column in date_columns or column in integer_columns else _EXCEL_LEFT
            if column in money_columns:
                cell.number_format = "#,##0.00"
            elif column in date_columns:
                cell.number_format = "yyyy-mm-dd"
            elif column in integer_columns:
                cell.number_format = "0"
    last_column = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A{start_row}:{last_column}{max(row_index, start_row)}"
    ws.freeze_panes = f"A{start_row + 1}"
    return row_index + 2


def _asset_values(row: ReportAssetRow) -> list[Any]:
    return [
        row.asset_no,
        row.name,
        row.category_name or "",
        row.model_name or "",
        row.serial_no or "",
        row.dept_name or "",
        row.custodian_name or "",
        row.status_label,
        row.purchase_date,
        row.purchase_price,
        row.warranty_end,
    ]


def _idle_values(row: IdleAssetRow) -> list[Any]:
    return [
        row.asset_no,
        row.name,
        row.category_name or "",
        row.dept_name or "",
        row.custodian_name or "",
        row.status_label,
        row.idle_since,
        row.idle_days,
        row.purchase_price,
        row.location or "",
    ]


def _compliance_values(row: ComplianceAssetRow) -> list[Any]:
    return [
        row.asset_no,
        row.name,
        row.category_name or "",
        row.dept_name or "",
        row.custodian_name or "",
        row.status_label,
        "、".join(item.label for item in row.issues),
        row.purchase_price,
        row.warranty_end,
        row.location or "",
    ]


_ASSET_HEADERS = [
    "资产编号",
    "资产名称",
    "分类",
    "型号",
    "序列号",
    "部门",
    "责任人",
    "状态",
    "购置日期",
    "购置金额（元）",
    "保修到期",
]
_ASSET_WIDTHS = [18, 24, 14, 20, 20, 16, 12, 10, 12, 14, 12]


def _metrics_rows(overview: AssetReportOverviewOut) -> list[tuple[str, Any]]:
    metrics = overview.metrics
    return [
        ("资产总数", metrics.total_count),
        ("资产总值（元）", round(metrics.total_value, 2)),
        ("平均单价（元）", round(metrics.avg_value, 2)),
        ("在用资产", metrics.in_use_count),
        ("在库资产", metrics.in_stock_count),
        ("闲置资产", metrics.idle_count),
        ("闲置资产价值（元）", round(metrics.idle_value, 2)),
        ("维修中资产", metrics.repairing_count),
        ("待报废资产", metrics.pending_scrap_count),
        ("已报废资产", metrics.scrapped_count),
        ("30 天内保修到期", metrics.warranty_expiring_30d),
        ("保修已到期", metrics.warranty_expired),
        ("存在合规问题的资产", metrics.compliance_asset_count),
    ]


def build_xlsx(
    *,
    report_type: str,
    title: str,
    filter_text: str,
    data_scope: str,
    generated_at: datetime,
    overview: AssetReportOverviewOut | None,
    rows: list[Any],
    compliance_items: list[ComplianceIssueSummary],
    summary_asset_count: int,
    today: date,
) -> bytes:
    workbook = Workbook()
    if overview is not None:
        metrics_sheet = workbook.active
        metrics_sheet.title = "概览指标"
        start = _write_intro(
            metrics_sheet, title, filter_text, data_scope, generated_at, 2
        )
        metrics_rows = _metrics_rows(overview)
        _write_table(
            metrics_sheet,
            ["指标", "数值"],
            metrics_rows,
            start_row=start,
            widths=[32, 24],
        )
        for row_index, (_, value) in enumerate(metrics_rows, start=start + 1):
            if isinstance(value, float):
                metrics_sheet.cell(row=row_index, column=2).number_format = "#,##0.00"

        status_sheet = workbook.create_sheet("状态分布")
        start = _write_intro(
            status_sheet, "状态分布", filter_text, data_scope, generated_at, 3
        )
        _write_table(
            status_sheet,
            ["状态", "数量", "资产价值（元）"],
            [
                (item.label, item.count, round(item.value, 2))
                for item in overview.status_distribution
            ],
            start_row=start,
            widths=[22, 14, 22],
            money_columns={3},
            integer_columns={2},
        )

        dept_sheet = workbook.create_sheet("部门分布")
        start = _write_intro(
            dept_sheet, "部门分布", filter_text, data_scope, generated_at, 3
        )
        _write_table(
            dept_sheet,
            ["部门", "数量", "资产价值（元）"],
            [
                (item.name, item.count, round(item.value, 2))
                for item in overview.dept_distribution
            ],
            start_row=start,
            widths=[32, 14, 22],
            money_columns={3},
            integer_columns={2},
        )

        category_sheet = workbook.create_sheet("分类分布")
        start = _write_intro(
            category_sheet, "分类分布", filter_text, data_scope, generated_at, 3
        )
        _write_table(
            category_sheet,
            ["分类", "数量", "资产价值（元）"],
            [
                (item.name, item.count, round(item.value, 2))
                for item in overview.category_distribution
            ],
            start_row=start,
            widths=[32, 14, 22],
            money_columns={3},
            integer_columns={2},
        )

        trend_sheet = workbook.create_sheet("采购趋势")
        start = _write_intro(
            trend_sheet, "近 12 个月采购趋势", filter_text, data_scope, generated_at, 3
        )
        _write_table(
            trend_sheet,
            ["月份", "新增数量", "采购金额（元）"],
            [
                (item.month, item.count, round(item.value, 2))
                for item in overview.purchase_trend
            ],
            start_row=start,
            widths=[18, 16, 22],
            money_columns={3},
            integer_columns={2},
        )
    else:
        sheet = workbook.active
        if report_type == "ASSETS":
            sheet.title = "资产清单"
            start = _write_intro(
                sheet, title, filter_text, data_scope, generated_at, len(_ASSET_HEADERS)
            )
            _write_table(
                sheet,
                _ASSET_HEADERS,
                [_asset_values(row) for row in rows],
                start_row=start,
                widths=_ASSET_WIDTHS,
                money_columns={10},
                date_columns={9, 11},
            )
        elif report_type == "IDLE":
            sheet.title = "闲置资产"
            headers = [
                "资产编号",
                "资产名称",
                "分类",
                "部门",
                "责任人",
                "状态",
                "闲置起始",
                "闲置天数",
                "购置金额（元）",
                "存放位置",
            ]
            start = _write_intro(
                sheet, title, filter_text, data_scope, generated_at, len(headers)
            )
            _write_table(
                sheet,
                headers,
                [_idle_values(row) for row in rows],
                start_row=start,
                widths=[18, 24, 14, 18, 12, 10, 18, 10, 14, 20],
                money_columns={9},
                date_columns={7},
                integer_columns={8},
            )
        else:
            sheet.title = "合规检查"
            headers = [
                "资产编号",
                "资产名称",
                "分类",
                "部门",
                "责任人",
                "状态",
                "问题项",
                "购置金额（元）",
                "保修到期",
                "存放位置",
            ]
            start = _write_intro(
                sheet, title, filter_text, data_scope, generated_at, len(headers)
            )
            _write_table(
                sheet,
                headers,
                [_compliance_values(row) for row in rows],
                start_row=start,
                widths=[18, 24, 14, 18, 12, 10, 30, 14, 12, 20],
                money_columns={8},
                date_columns={9},
            )
            summary_start = start + len(rows) + 4
            sheet.cell(
                row=summary_start,
                column=1,
                value="合规问题汇总",
            ).font = Font(size=11, bold=True, color="1F4E78")
            _write_table(
                sheet,
                ["问题项", "数量", "处理建议"],
                [
                    (item.label, item.count, item.suggestion)
                    for item in compliance_items
                ],
                start_row=summary_start + 1,
                widths=[18, 10, 60],
                integer_columns={2},
            )
            sheet.cell(
                row=summary_start + 1,
                column=1,
                value=(
                    f"存在合规问题的资产共 {summary_asset_count} 台"
                    f"（统计日期 {today.isoformat()}）"
                ),
            )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class _NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs) -> None:
        super().__init__(*args, **kwargs)
        self._saved_page_states: list[dict[str, Any]] = []

    def showPage(self) -> None:
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self) -> None:
        page_count = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.setFont(_PDF_FONT, 8)
            self.drawRightString(
                self._pagesize[0] - 30,
                20,
                f"第 {self._pageNumber} / {page_count} 页",
            )
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)


def _pdf_table(
    headers: list[str],
    rows: list[list[Any]],
    widths: list[float],
    *,
    center_columns: set[int] | None = None,
) -> Table:
    center_columns = center_columns or set()
    header_cells = [Paragraph(escape(header), _PDF_HEADER_CELL) for header in headers]
    body = []
    for values in rows:
        cells = []
        for index, value in enumerate(values, start=1):
            style = _PDF_CELL_CENTER if index in center_columns else _PDF_CELL
            cells.append(Paragraph(escape("" if value is None else str(value)), style))
        body.append(cells)
    table = Table(
        [header_cells, *body],
        colWidths=widths,
        repeatRows=1,
        splitByRow=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D3DF")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7FB")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def _pdf_section(
    story: list[Any],
    heading: str,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[float],
    *,
    center_columns: set[int] | None = None,
) -> None:
    story.append(Paragraph(escape(heading), _PDF_HEADING))
    if not rows:
        story.append(Paragraph("无符合条件的数据。", _PDF_META))
        return
    story.append(
        _pdf_table(headers, rows, widths, center_columns=center_columns)
    )


def build_pdf(
    *,
    report_type: str,
    title: str,
    filter_text: str,
    data_scope: str,
    generated_at: datetime,
    overview: AssetReportOverviewOut | None,
    rows: list[Any],
    compliance_items: list[ComplianceIssueSummary],
    summary_asset_count: int,
    today: date,
    row_limit: int,
) -> bytes:
    buffer = BytesIO()
    is_landscape = report_type in ("ASSETS", "IDLE", "COMPLIANCE")
    page_size = landscape(A4) if is_landscape else A4
    document = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=13 * mm,
        bottomMargin=15 * mm,
        title=f"ITAM {title}",
        author="ITAM",
    )
    story: list[Any] = [
        Paragraph(f"ITAM {escape(title)}", _PDF_TITLE),
        Paragraph(f"筛选条件：{escape(filter_text)}", _PDF_META),
        Paragraph(
            f"数据范围：{escape(data_scope)}　"
            f"生成时间：{generated_at.strftime('%Y-%m-%d %H:%M:%S')}",
            _PDF_META,
        ),
        Spacer(1, 4),
    ]

    if overview is not None:
        _pdf_section(
            story,
            "核心指标",
            ["指标", "数值"],
            [
                [name, _fmt_money(value) if isinstance(value, float) else str(value)]
                for name, value in _metrics_rows(overview)
            ],
            [250, 230],
            center_columns={2},
        )
        _pdf_section(
            story,
            "状态分布",
            ["状态", "数量", "资产价值（元）"],
            [
                [item.label, str(item.count), _fmt_money(item.value)]
                for item in overview.status_distribution
            ],
            [190, 120, 170],
            center_columns={2, 3},
        )
        _pdf_section(
            story,
            "部门分布",
            ["部门", "数量", "资产价值（元）"],
            [
                [item.name, str(item.count), _fmt_money(item.value)]
                for item in overview.dept_distribution
            ],
            [260, 100, 120],
            center_columns={2, 3},
        )
        _pdf_section(
            story,
            "分类分布",
            ["分类", "数量", "资产价值（元）"],
            [
                [item.name, str(item.count), _fmt_money(item.value)]
                for item in overview.category_distribution
            ],
            [260, 100, 120],
            center_columns={2, 3},
        )
        _pdf_section(
            story,
            "近 12 个月采购趋势",
            ["月份", "新增数量", "采购金额（元）"],
            [
                [item.month, str(item.count), _fmt_money(item.value)]
                for item in overview.purchase_trend
            ],
            [190, 120, 170],
            center_columns={2, 3},
        )
    else:
        total = len(rows)
        display_rows = rows[:row_limit]
        if total > row_limit:
            story.append(
                Paragraph(
                    f"共 {total} 条数据，本 PDF 仅展示前 {row_limit} 条，"
                    "完整数据请使用 Excel 导出。",
                    _PDF_META,
                )
            )
        if report_type == "ASSETS":
            _pdf_section(
                story,
                "资产清单",
                ["资产编号", "资产名称", "分类", "部门", "责任人", "状态", "购置日期", "金额（元）", "保修到期", "位置"],
                [
                    [
                        row.asset_no,
                        row.name,
                        row.category_name or "",
                        row.dept_name or "",
                        row.custodian_name or "",
                        row.status_label,
                        _fmt_date(row.purchase_date),
                        _fmt_money(row.purchase_price),
                        _fmt_date(row.warranty_end),
                        row.location or "",
                    ]
                    for row in display_rows
                ],
                [86, 105, 70, 70, 48, 42, 58, 60, 58, 90],
                center_columns={7, 8, 9},
            )
        elif report_type == "IDLE":
            _pdf_section(
                story,
                "闲置资产",
                ["资产编号", "资产名称", "分类", "部门", "责任人", "状态", "闲置起始", "闲置天数", "金额（元）", "位置"],
                [
                    [
                        row.asset_no,
                        row.name,
                        row.category_name or "",
                        row.dept_name or "",
                        row.custodian_name or "",
                        row.status_label,
                        _fmt_date(row.idle_since),
                        f"{row.idle_days} 天",
                        _fmt_money(row.purchase_price),
                        row.location or "",
                    ]
                    for row in display_rows
                ],
                [86, 105, 66, 66, 46, 42, 62, 42, 58, 86],
                center_columns={7, 8, 9},
            )
        else:
            _pdf_section(
                story,
                "合规问题汇总",
                ["问题项", "数量", "处理建议"],
                [
                    [item.label, str(item.count), item.suggestion]
                    for item in compliance_items
                ],
                [110, 55, 580],
                center_columns={2},
            )
            story.append(
                Paragraph(
                    f"存在合规问题的资产共 {summary_asset_count} 台"
                    f"（统计日期 {today.isoformat()}）",
                    _PDF_META,
                )
            )
            _pdf_section(
                story,
                "合规问题明细",
                ["资产编号", "资产名称", "分类", "部门", "状态", "问题项", "金额（元）", "保修到期", "位置"],
                [
                    [
                        row.asset_no,
                        row.name,
                        row.category_name or "",
                        row.dept_name or "",
                        row.status_label,
                        "、".join(item.label for item in row.issues),
                        _fmt_money(row.purchase_price),
                        _fmt_date(row.warranty_end),
                        row.location or "",
                    ]
                    for row in display_rows
                ],
                [82, 100, 62, 62, 42, 150, 56, 56, 80],
                center_columns={7, 8, 9},
            )

    document.build(story, canvasmaker=_NumberedCanvas)
    return buffer.getvalue()
