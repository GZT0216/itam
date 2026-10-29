from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import case, false, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.asset_ops import (
    INVENTORY_ADJUSTMENT_NO_PREFIX,
    INVENTORY_ADJUSTMENT_STATUS_LABELS,
    INVENTORY_ADJUSTMENT_STATUS_TAG_TYPES,
    INVENTORY_NO_PREFIX,
    INVENTORY_RESULT_LABELS,
    INVENTORY_RESULT_TAG_TYPES,
    INVENTORY_STATUS_LABELS,
    INVENTORY_STATUS_TAG_TYPES,
    InventoryAdjustmentStatus,
    InventoryResult,
    InventoryStatus,
)
from app.core.assets import ASSET_STATUS_LABELS, AssetStatus
from app.core.exceptions import BizError, ConflictError, NotFoundError, PermissionDenied
from app.models import (
    Asset,
    AssetCategory,
    AssetHistory,
    AssetModel,
    Dept,
    InventoryAdjustment,
    InventoryItem,
    InventorySurplus,
    InventoryTask,
)
from app.models.base import now_cst
from app.schemas.asset_ops import (
    InventoryAdjustmentCreate,
    InventoryAdjustmentOut,
    InventoryAdjustmentReview,
    InventoryCancelPayload,
    InventoryCreate,
    InventoryItemOut,
    InventoryItemPageOut,
    InventoryItemUpdate,
    InventoryReportOut,
    InventoryScanPayload,
    InventorySurplusCreate,
    InventorySurplusOut,
    InventoryTaskOut,
    InventoryTaskPageOut,
)
from app.services.audit_service import AuditService
from app.services.asset_snapshot import asset_snapshot
from app.services.permission_service import AuthContext
from app.services.scope_service import scoped_dept_ids
from app.services.txn import commit_or_conflict

STATUS_CONFLICT_MESSAGE = "资产状态已变更，请刷新后重试"
MANAGE_PERMISSION = "asset:inventory:manage"
APPROVE_PERMISSION = "asset:inventory:approve"


@dataclass
class InventoryReportExport:
    content: bytes
    filename: str
    media_type: str
    row_count: int


class InventoryService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.audit = AuditService(session)

    async def _next_no(self) -> str:
        head = f"{INVENTORY_NO_PREFIX}-{now_cst():%Y%m}-"
        max_no = await self.session.scalar(
            select(func.max(InventoryTask.task_no)).where(
                InventoryTask.task_no.like(f"{head}%")
            )
        )
        sequence = 1
        if max_no:
            suffix = max_no[len(head) :]
            if suffix.isdigit():
                sequence = int(suffix) + 1
        return f"{head}{sequence:04d}"

    async def _allowed_depts(self, context: AuthContext) -> set[int] | None:
        return await scoped_dept_ids(self.session, context)

    @staticmethod
    def _in_scope(task: InventoryTask, allowed: set[int] | None) -> bool:
        if allowed is None:
            return True
        return task.scope_dept_id is not None and task.scope_dept_id in allowed

    async def _get_task(self, task_id: int) -> InventoryTask:
        task = await self.session.get(InventoryTask, task_id)
        if task is None:
            raise NotFoundError("盘点任务不存在")
        return task

    async def _get_task_for_view(
        self, task_id: int, context: AuthContext
    ) -> tuple[InventoryTask, set[int] | None]:
        task = await self._get_task(task_id)
        allowed = await self._allowed_depts(context)
        if not self._in_scope(task, allowed):
            raise PermissionDenied("无权查看该盘点任务")
        return task, allowed

    async def _get_task_for_manage(
        self, task_id: int, context: AuthContext
    ) -> tuple[InventoryTask, set[int] | None]:
        task, allowed = await self._get_task_for_view(task_id, context)
        if not context.has_permission(MANAGE_PERMISSION):
            raise PermissionDenied("无权执行盘点操作")
        if task.status != InventoryStatus.IN_PROGRESS:
            raise BizError("仅进行中的盘点任务可以操作")
        return task, allowed

    async def _resolve_dept_ids(
        self,
        scope_dept_id: int | None,
        include_child: bool,
        allowed: set[int] | None,
    ) -> set[int] | None:
        if scope_dept_id is None:
            return allowed
        dept = await self.session.get(Dept, scope_dept_id)
        if dept is None:
            raise BizError("盘点范围部门不存在")
        if allowed is not None and scope_dept_id not in allowed:
            raise PermissionDenied("无权在该部门范围内创建盘点任务")
        ids = {dept.id}
        if include_child:
            rows = (
                await self.session.execute(select(Dept.id, Dept.path))
            ).all()
            ids |= {
                row.id
                for row in rows
                if row.id != dept.id and row.path.startswith(dept.path)
            }
        if allowed is not None:
            ids &= allowed
        return ids

    async def _resolve_category_ids(
        self, scope_category_id: int | None, include_child: bool
    ) -> set[int] | None:
        if scope_category_id is None:
            return None
        category = await self.session.get(AssetCategory, scope_category_id)
        if category is None:
            raise BizError("盘点范围分类不存在")
        ids = {category.id}
        if include_child:
            rows = (
                await self.session.execute(select(AssetCategory.id, AssetCategory.path))
            ).all()
            ids |= {
                row.id
                for row in rows
                if row.id != category.id and row.path.startswith(category.path)
            }
        return ids

    async def _stats(self, task_ids: list[int]) -> dict[int, dict[str, int]]:
        if not task_ids:
            return {}
        rows = (
            await self.session.execute(
                select(
                    InventoryItem.task_id,
                    func.count().label("total"),
                    func.sum(
                        case((InventoryItem.result != InventoryResult.PENDING, 1), else_=0)
                    ).label("checked"),
                    func.sum(
                        case((InventoryItem.result == InventoryResult.NORMAL, 1), else_=0)
                    ).label("normal"),
                    func.sum(
                        case((InventoryItem.result == InventoryResult.ABNORMAL, 1), else_=0)
                    ).label("abnormal"),
                    func.sum(
                        case((InventoryItem.result == InventoryResult.MISSING, 1), else_=0)
                    ).label("missing"),
                )
                .where(InventoryItem.task_id.in_(task_ids))
                .group_by(InventoryItem.task_id)
            )
        ).all()
        surplus_rows = (
            await self.session.execute(
                select(InventorySurplus.task_id, func.count())
                .where(InventorySurplus.task_id.in_(task_ids))
                .group_by(InventorySurplus.task_id)
            )
        ).all()
        surplus_counts = {task_id: int(count) for task_id, count in surplus_rows}
        result: dict[int, dict[str, int]] = {}
        for row in rows:
            result[row.task_id] = {
                "total": int(row.total or 0),
                "checked": int(row.checked or 0),
                "normal": int(row.normal or 0),
                "abnormal": int(row.abnormal or 0),
                "missing": int(row.missing or 0),
                "surplus": surplus_counts.pop(row.task_id, 0),
                "pending": int(row.total or 0) - int(row.checked or 0),
            }
        for task_id, count in surplus_counts.items():
            result[task_id] = {
                "total": 0,
                "checked": 0,
                "normal": 0,
                "abnormal": 0,
                "missing": 0,
                "surplus": count,
                "pending": 0,
            }
        return result

    async def _latest_adjustments(
        self, task_ids: list[int]
    ) -> dict[int, InventoryAdjustment]:
        if not task_ids:
            return {}
        rows = (
            (
                await self.session.execute(
                    select(InventoryAdjustment)
                    .where(InventoryAdjustment.task_id.in_(task_ids))
                    .order_by(InventoryAdjustment.id.desc())
                )
            )
            .scalars()
            .all()
        )
        result: dict[int, InventoryAdjustment] = {}
        for row in rows:
            result.setdefault(row.task_id, row)
        return result

    def _task_out(
        self,
        task: InventoryTask,
        context: AuthContext,
        allowed: set[int] | None,
        stats: dict[str, int] | None = None,
        adjustment: InventoryAdjustment | None = None,
    ) -> InventoryTaskOut:
        counts = stats or {}
        in_scope = self._in_scope(task, allowed)
        missing_count = counts.get("missing", 0)
        surplus_count = counts.get("surplus", 0)
        has_pending = bool(
            adjustment and adjustment.status == InventoryAdjustmentStatus.PENDING
        )
        can_submit = bool(
            task.status == InventoryStatus.COMPLETED
            and in_scope
            and context.has_permission(MANAGE_PERMISSION)
            and (missing_count > 0 or surplus_count > 0)
            and (
                adjustment is None
                or adjustment.status == InventoryAdjustmentStatus.REJECTED
            )
        )
        return InventoryTaskOut(
            id=task.id,
            task_no=task.task_no,
            name=task.name,
            status=task.status,
            status_label=INVENTORY_STATUS_LABELS.get(task.status, task.status),
            status_tag_type=INVENTORY_STATUS_TAG_TYPES.get(task.status, "info"),
            scope_dept_id=task.scope_dept_id,
            scope_dept_name=task.scope_dept_name,
            include_child=task.include_child,
            scope_category_id=task.scope_category_id,
            scope_category_name=task.scope_category_name,
            include_child_category=task.include_child_category,
            planned_date=task.planned_date,
            remark=task.remark,
            started_at=task.started_at,
            completed_at=task.completed_at,
            cancelled_at=task.cancelled_at,
            cancel_reason=task.cancel_reason,
            created_by_name=task.created_by_name,
            total_assets=counts.get("total", 0),
            checked_assets=counts.get("checked", 0),
            normal_count=counts.get("normal", 0),
            abnormal_count=counts.get("abnormal", 0),
            missing_count=missing_count,
            surplus_count=surplus_count,
            pending_count=counts.get("pending", 0),
            can_manage=(
                task.status == InventoryStatus.IN_PROGRESS
                and context.has_permission(MANAGE_PERMISSION)
                and in_scope
            ),
            adjustment_id=adjustment.id if adjustment else None,
            adjustment_no=adjustment.adjustment_no if adjustment else None,
            adjustment_status=adjustment.status if adjustment else None,
            adjustment_status_label=(
                INVENTORY_ADJUSTMENT_STATUS_LABELS.get(adjustment.status)
                if adjustment
                else None
            ),
            adjustment_status_tag_type=(
                INVENTORY_ADJUSTMENT_STATUS_TAG_TYPES.get(adjustment.status)
                if adjustment
                else None
            ),
            adjustment_version=adjustment.version if adjustment else None,
            can_submit_adjustment=can_submit,
            can_approve_adjustment=bool(
                has_pending
                and in_scope
                and context.has_permission(APPROVE_PERMISSION)
            ),
            version=task.version,
            created_at=task.created_at,
            updated_at=task.updated_at,
        )

    @staticmethod
    def _item_out(item: InventoryItem) -> InventoryItemOut:
        return InventoryItemOut(
            id=item.id,
            task_id=item.task_id,
            asset_id=item.asset_id,
            asset_no=item.asset_no,
            asset_name=item.asset_name,
            category_name=item.category_name,
            dept_name=item.dept_name,
            location=item.location,
            custodian_name=item.custodian_name,
            expected_status=item.expected_status,
            expected_status_label=ASSET_STATUS_LABELS.get(
                item.expected_status, item.expected_status
            ),
            result=item.result,
            result_label=INVENTORY_RESULT_LABELS.get(item.result, item.result),
            result_tag_type=INVENTORY_RESULT_TAG_TYPES.get(item.result, "info"),
            actual_location=item.actual_location,
            remark=item.remark,
            checked_by_name=item.checked_by_name,
            checked_at=item.checked_at,
            created_at=item.created_at,
        )

    @staticmethod
    def _surplus_out(row: InventorySurplus) -> InventorySurplusOut:
        return InventorySurplusOut(
            id=row.id,
            task_id=row.task_id,
            asset_no=row.asset_no,
            name=row.name,
            category_id=row.category_id,
            category_name=row.category_name,
            model_id=row.model_id,
            model_name=row.model.name if row.model else None,
            dept_id=row.dept_id,
            dept_name=row.dept_name,
            serial_no=row.serial_no,
            location=row.location,
            purchase_date=row.purchase_date,
            purchase_price=(
                float(row.purchase_price) if row.purchase_price is not None else None
            ),
            remark=row.remark,
            created_by_name=row.created_by_name,
            created_asset_id=row.created_asset_id,
            applied_at=row.applied_at,
            created_at=row.created_at,
        )

    @staticmethod
    def _adjustment_out(row: InventoryAdjustment) -> InventoryAdjustmentOut:
        return InventoryAdjustmentOut(
            id=row.id,
            task_id=row.task_id,
            task_no=row.task_no,
            task_name=row.task_name,
            adjustment_no=row.adjustment_no,
            status=row.status,
            status_label=INVENTORY_ADJUSTMENT_STATUS_LABELS.get(row.status, row.status),
            status_tag_type=INVENTORY_ADJUSTMENT_STATUS_TAG_TYPES.get(
                row.status, "info"
            ),
            loss_count=row.loss_count,
            surplus_count=row.surplus_count,
            reason=row.reason,
            requested_by_name=row.requested_by_name,
            requested_at=row.requested_at,
            reviewed_by_name=row.reviewed_by_name,
            reviewed_at=row.reviewed_at,
            review_comment=row.review_comment,
            version=row.version,
            created_at=row.created_at,
            updated_at=row.updated_at,
        )

    async def _next_adjustment_no(self) -> str:
        head = f"{INVENTORY_ADJUSTMENT_NO_PREFIX}-{now_cst():%Y%m}-"
        max_no = await self.session.scalar(
            select(func.max(InventoryAdjustment.adjustment_no)).where(
                InventoryAdjustment.adjustment_no.like(f"{head}%")
            )
        )
        sequence = 1
        if max_no:
            suffix = max_no[len(head) :]
            if suffix.isdigit():
                sequence = int(suffix) + 1
        return f"{head}{sequence:04d}"

    async def _latest_adjustment(
        self, task_id: int
    ) -> InventoryAdjustment | None:
        return (
            await self.session.execute(
                select(InventoryAdjustment)
                .where(InventoryAdjustment.task_id == task_id)
                .order_by(InventoryAdjustment.id.desc())
                .limit(1)
            )
        ).scalar_one_or_none()

    async def _assert_surplus_code_available(
        self, asset_no: str, *, exclude_id: int | None = None
    ) -> None:
        code = asset_no.strip()
        asset_exists = await self.session.scalar(
            select(Asset.id).where(func.lower(Asset.asset_no) == code.casefold())
        )
        if asset_exists is not None:
            raise BizError(f"资产编号「{code}」已存在，请核对后重新登记")
        surplus_stmt = select(InventorySurplus.id).where(
            func.lower(InventorySurplus.asset_no) == code.casefold(),
            InventorySurplus.applied_at.is_(None),
        )
        if exclude_id is not None:
            surplus_stmt = surplus_stmt.where(InventorySurplus.id != exclude_id)
        if await self.session.scalar(surplus_stmt) is not None:
            raise BizError(f"资产编号「{code}」已在其他盘盈记录中登记")

    async def list_surplus(
        self, task_id: int, context: AuthContext
    ) -> list[InventorySurplusOut]:
        await self._get_task_for_view(task_id, context)
        rows = (
            (
                await self.session.execute(
                    select(InventorySurplus)
                    .where(InventorySurplus.task_id == task_id)
                    .order_by(InventorySurplus.id)
                )
            )
            .scalars()
            .all()
        )
        return [self._surplus_out(row) for row in rows]

    async def add_surplus(
        self,
        task_id: int,
        payload: InventorySurplusCreate,
        context: AuthContext,
    ) -> InventorySurplusOut:
        task, allowed = await self._get_task_for_view(task_id, context)
        if not context.has_permission(MANAGE_PERMISSION):
            raise PermissionDenied("无权执行盘点操作")
        if task.status != InventoryStatus.IN_PROGRESS:
            raise BizError("仅盘点中的任务可以登记盘盈资产")
        category = await self.session.get(AssetCategory, payload.category_id)
        if category is None or category.status != 1:
            raise BizError("盘盈资产分类不存在或已停用")
        child_count = int(
            await self.session.scalar(
                select(func.count())
                .select_from(AssetCategory)
                .where(
                    AssetCategory.parent_id == category.id,
                    AssetCategory.status == 1,
                )
            )
            or 0
        )
        if child_count:
            raise BizError("请选择最终一级分类登记盘盈资产")
        model = None
        if payload.model_id is not None:
            model = await self.session.get(AssetModel, payload.model_id)
            if model is None or model.status != 1:
                raise BizError("盘盈资产型号不存在或已停用")
            if model.category_id != category.id:
                raise BizError("盘盈资产型号与分类不一致")
        dept_id = payload.dept_id if payload.dept_id is not None else task.scope_dept_id
        dept = None
        if dept_id is not None:
            dept = await self.session.get(Dept, dept_id)
            if dept is None:
                raise BizError("盘盈资产所属部门不存在")
            if allowed is not None and dept.id not in allowed:
                raise PermissionDenied("无权在该部门范围内登记盘盈资产")
        else:
            dept_id = context.user.dept_id
            dept = context.user.dept
        await self._assert_surplus_code_available(payload.asset_no)
        row = InventorySurplus(
            task_id=task.id,
            asset_no=payload.asset_no.strip(),
            name=payload.name,
            category=category,
            category_id=category.id,
            category_name=category.name,
            model=model,
            model_id=model.id if model else None,
            dept_id=dept_id,
            dept_name=dept.name if dept else None,
            serial_no=payload.serial_no,
            location=payload.location,
            purchase_date=payload.purchase_date,
            purchase_price=(
                Decimal(str(payload.purchase_price))
                if payload.purchase_price is not None
                else None
            ),
            remark=payload.remark,
            created_by_id=context.user.id,
            created_by_name=context.user.real_name or context.user.username,
        )
        self.session.add(row)
        await self.session.flush()
        await self.audit.record(
            action="asset.inventory.surplus.create",
            module="资产盘点",
            summary=f"登记盘盈资产：{row.asset_no} {row.name}",
            actor=context.user,
            resource_type="inventory_task",
            resource_id=task.id,
            after={
                "task_no": task.task_no,
                "asset_no": row.asset_no,
                "name": row.name,
                "category_name": row.category_name,
                "dept_name": row.dept_name,
            },
        )
        await commit_or_conflict(
            self.session, message=STATUS_CONFLICT_MESSAGE, code=40911
        )
        return self._surplus_out(row)

    async def delete_surplus(
        self, task_id: int, surplus_id: int, context: AuthContext
    ) -> None:
        task, _ = await self._get_task_for_view(task_id, context)
        if not context.has_permission(MANAGE_PERMISSION):
            raise PermissionDenied("无权执行盘点操作")
        if task.status != InventoryStatus.IN_PROGRESS:
            raise BizError("仅盘点中的任务可以调整盘盈登记")
        row = await self.session.get(InventorySurplus, surplus_id)
        if row is None or row.task_id != task.id:
            raise NotFoundError("盘盈记录不存在")
        if row.applied_at is not None:
            raise BizError("盘盈资产已核销，无法删除")
        await self.audit.record(
            action="asset.inventory.surplus.delete",
            module="资产盘点",
            summary=f"删除盘盈登记：{row.asset_no} {row.name}",
            actor=context.user,
            resource_type="inventory_task",
            resource_id=task.id,
            before={
                "asset_no": row.asset_no,
                "name": row.name,
                "category_name": row.category_name,
            },
        )
        await self.session.delete(row)
        await commit_or_conflict(
            self.session, message=STATUS_CONFLICT_MESSAGE, code=40911
        )

    async def submit_adjustment(
        self,
        task_id: int,
        payload: InventoryAdjustmentCreate,
        context: AuthContext,
    ) -> InventoryAdjustmentOut:
        task, _ = await self._get_task_for_view(task_id, context)
        if not context.has_permission(MANAGE_PERMISSION):
            raise PermissionDenied("无权提交盘点处理")
        if task.status != InventoryStatus.COMPLETED:
            raise BizError("请先完成盘点后再提交盘盈盘亏处理")
        latest = await self._latest_adjustment(task.id)
        if latest is not None and latest.status in (
            InventoryAdjustmentStatus.PENDING,
            InventoryAdjustmentStatus.APPROVED,
        ):
            raise BizError("该盘点任务已有待审批或已通过的处理单")
        stats = (await self._stats([task.id])).get(task.id, {})
        loss_count = int(stats.get("missing", 0))
        surplus_count = int(stats.get("surplus", 0))
        if loss_count == 0 and surplus_count == 0:
            raise BizError("当前盘点没有盘盈盘亏，无需提交处理")
        row = InventoryAdjustment(
            task_id=task.id,
            task_no=task.task_no,
            task_name=task.name,
            adjustment_no=await self._next_adjustment_no(),
            status=InventoryAdjustmentStatus.PENDING,
            loss_count=loss_count,
            surplus_count=surplus_count,
            reason=payload.reason,
            requested_by_id=context.user.id,
            requested_by_name=context.user.real_name or context.user.username,
            requested_at=now_cst(),
        )
        self.session.add(row)
        await self.session.flush()
        await self.audit.record(
            action="asset.inventory.adjustment.submit",
            module="资产盘点",
            summary=(
                f"提交盘点盈亏处理：{task.task_no}，"
                f"盘亏 {loss_count} 项，盘盈 {surplus_count} 项"
            ),
            actor=context.user,
            resource_type="inventory_task",
            resource_id=task.id,
            after={
                "adjustment_no": row.adjustment_no,
                "status": row.status,
                "loss_count": loss_count,
                "surplus_count": surplus_count,
                "reason": payload.reason,
            },
        )
        await commit_or_conflict(
            self.session, message=STATUS_CONFLICT_MESSAGE, code=40911
        )
        return self._adjustment_out(row)

    async def get_adjustment(
        self, task_id: int, context: AuthContext
    ) -> InventoryAdjustmentOut | None:
        await self._get_task_for_view(task_id, context)
        latest = await self._latest_adjustment(task_id)
        return self._adjustment_out(latest) if latest else None

    async def _apply_adjustment(
        self,
        task: InventoryTask,
        adjustment: InventoryAdjustment,
        context: AuthContext,
    ) -> None:
        missing_items = list(
            (
                await self.session.execute(
                    select(InventoryItem).where(
                        InventoryItem.task_id == task.id,
                        InventoryItem.result == InventoryResult.MISSING,
                    )
                )
            )
            .scalars()
            .all()
        )
        asset_ids = [item.asset_id for item in missing_items]
        assets = (
            list(
                (
                    await self.session.execute(
                        select(Asset)
                        .where(Asset.id.in_(asset_ids))
                        .with_for_update()
                    )
                )
                .scalars()
                .all()
            )
            if asset_ids
            else []
        )
        asset_map = {asset.id: asset for asset in assets}
        for item in missing_items:
            asset = asset_map.get(item.asset_id)
            if asset is None:
                raise BizError(f"盘亏资产「{item.asset_no}」不存在")
            if asset.status in (AssetStatus.SCRAPPED, AssetStatus.PENDING_STOCK):
                raise BizError(
                    f"资产「{item.asset_no}」状态已变更，请刷新后重新提交处理"
                )
            before = asset_snapshot(asset)
            asset.status = AssetStatus.SCRAPPED
            asset.custodian_id = None
            self.session.add(
                AssetHistory(
                    asset_id=asset.id,
                    action="INVENTORY_LOSS",
                    action_label="盘点盘亏核销",
                    from_status=before["status"],
                    to_status=asset.status,
                    before_snapshot=before,
                    after_snapshot=asset_snapshot(asset),
                    operator_id=context.user.id,
                    operator_name=context.user.real_name or context.user.username,
                    summary=f"盘点任务 {task.task_no} 盘亏核销",
                )
            )

        surplus_rows = list(
            (
                await self.session.execute(
                    select(InventorySurplus)
                    .where(
                        InventorySurplus.task_id == task.id,
                        InventorySurplus.applied_at.is_(None),
                    )
                    .order_by(InventorySurplus.id)
                    .with_for_update()
                )
            )
            .scalars()
            .all()
        )
        created_assets: list[str] = []
        now = now_cst()
        for row in surplus_rows:
            await self._assert_surplus_code_available(
                row.asset_no, exclude_id=row.id
            )
            category = await self.session.get(AssetCategory, row.category_id)
            if category is None:
                raise BizError(f"盘盈资产「{row.asset_no}」分类不存在")
            asset = Asset(
                asset_no=row.asset_no,
                name=row.name,
                category_id=row.category_id,
                category=category,
                model_id=row.model_id,
                serial_no=row.serial_no,
                location=row.location,
                dept_id=row.dept_id,
                purchase_date=row.purchase_date,
                purchase_price=row.purchase_price,
                status=AssetStatus.IN_STOCK,
                remark=row.remark or f"盘点盘盈（{task.task_no}）",
            )
            self.session.add(asset)
            await self.session.flush()
            row.applied_at = now
            row.created_asset_id = asset.id
            created_assets.append(asset.asset_no)

        adjustment.status = InventoryAdjustmentStatus.APPROVED
        adjustment.reviewed_by_id = context.user.id
        adjustment.reviewed_by_name = context.user.real_name or context.user.username
        adjustment.reviewed_at = now
        await self.audit.record(
            action="asset.inventory.adjustment.approve",
            module="资产盘点",
            summary=(
                f"盘点盈亏处理审批通过：{adjustment.adjustment_no}，"
                f"核销盘亏 {len(missing_items)} 项，盘盈建档 {len(created_assets)} 项"
            ),
            actor=context.user,
            resource_type="inventory_task",
            resource_id=task.id,
            before={"status": InventoryAdjustmentStatus.PENDING},
            after={
                "adjustment_no": adjustment.adjustment_no,
                "status": adjustment.status,
                "loss_assets": [item.asset_no for item in missing_items],
                "surplus_assets": created_assets,
            },
        )

    async def approve_adjustment(
        self,
        task_id: int,
        payload: InventoryAdjustmentReview,
        context: AuthContext,
    ) -> InventoryAdjustmentOut:
        task, allowed = await self._get_task_for_view(task_id, context)
        if not context.has_permission(APPROVE_PERMISSION):
            raise PermissionDenied("无权审批盘点盈亏处理")
        if not self._in_scope(task, allowed):
            raise PermissionDenied("无权审批该盘点任务")
        adjustment = await self._latest_adjustment(task.id)
        if adjustment is None or adjustment.status != InventoryAdjustmentStatus.PENDING:
            raise BizError("当前盘点任务没有待审批的盈亏处理单")
        if payload.version != adjustment.version:
            raise ConflictError("处理单已被其他人更新，请刷新后重试", code=40912)
        await self._apply_adjustment(task, adjustment, context)
        adjustment.review_comment = payload.comment
        await commit_or_conflict(
            self.session, message=STATUS_CONFLICT_MESSAGE, code=40911
        )
        return self._adjustment_out(adjustment)

    async def reject_adjustment(
        self,
        task_id: int,
        payload: InventoryAdjustmentReview,
        context: AuthContext,
    ) -> InventoryAdjustmentOut:
        task, allowed = await self._get_task_for_view(task_id, context)
        if not context.has_permission(APPROVE_PERMISSION):
            raise PermissionDenied("无权审批盘点盈亏处理")
        if not self._in_scope(task, allowed):
            raise PermissionDenied("无权审批该盘点任务")
        if not payload.comment:
            raise BizError("请填写驳回原因")
        adjustment = await self._latest_adjustment(task.id)
        if adjustment is None or adjustment.status != InventoryAdjustmentStatus.PENDING:
            raise BizError("当前盘点任务没有待审批的盈亏处理单")
        if payload.version != adjustment.version:
            raise ConflictError("处理单已被其他人更新，请刷新后重试", code=40912)
        adjustment.status = InventoryAdjustmentStatus.REJECTED
        adjustment.reviewed_by_id = context.user.id
        adjustment.reviewed_by_name = context.user.real_name or context.user.username
        adjustment.reviewed_at = now_cst()
        adjustment.review_comment = payload.comment
        await self.audit.record(
            action="asset.inventory.adjustment.reject",
            module="资产盘点",
            summary=f"驳回盘点盈亏处理：{adjustment.adjustment_no}",
            actor=context.user,
            resource_type="inventory_task",
            resource_id=task.id,
            before={"status": InventoryAdjustmentStatus.PENDING},
            after={
                "adjustment_no": adjustment.adjustment_no,
                "status": adjustment.status,
                "reason": payload.comment,
            },
        )
        await commit_or_conflict(
            self.session, message=STATUS_CONFLICT_MESSAGE, code=40911
        )
        return self._adjustment_out(adjustment)

    async def _report_data(
        self, task_id: int, context: AuthContext
    ) -> InventoryReportOut:
        task, allowed = await self._get_task_for_view(task_id, context)
        stats = (await self._stats([task.id])).get(task.id, {})
        adjustment = await self._latest_adjustment(task.id)
        items = (
            (
                await self.session.execute(
                    select(InventoryItem)
                    .where(
                        InventoryItem.task_id == task.id,
                        InventoryItem.result.in_(
                            (InventoryResult.MISSING, InventoryResult.ABNORMAL)
                        ),
                    )
                    .order_by(InventoryItem.id)
                )
            )
            .scalars()
            .all()
        )
        surplus_rows = (
            (
                await self.session.execute(
                    select(InventorySurplus)
                    .where(InventorySurplus.task_id == task.id)
                    .order_by(InventorySurplus.id)
                )
            )
            .scalars()
            .all()
        )
        return InventoryReportOut(
            task=self._task_out(task, context, allowed, stats, adjustment),
            missing_items=[
                self._item_out(row)
                for row in items
                if row.result == InventoryResult.MISSING
            ],
            abnormal_items=[
                self._item_out(row)
                for row in items
                if row.result == InventoryResult.ABNORMAL
            ],
            surplus_items=[self._surplus_out(row) for row in surplus_rows],
            adjustment=(
                self._adjustment_out(adjustment) if adjustment is not None else None
            ),
            generated_at=now_cst(),
        )

    async def report(self, task_id: int, context: AuthContext) -> InventoryReportOut:
        return await self._report_data(task_id, context)

    @staticmethod
    def _style_sheet(sheet, widths: list[int]) -> None:
        header_fill = PatternFill("solid", fgColor="DCE6F1")
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

    async def export_report(
        self, task_id: int, context: AuthContext
    ) -> InventoryReportExport:
        data = await self._report_data(task_id, context)
        workbook = Workbook()
        overview = workbook.active
        overview.title = "盘点概览"
        overview.append(["项目", "内容"])
        overview.append(["盘点单号", data.task.task_no])
        overview.append(["盘点名称", data.task.name])
        overview.append(["盘点状态", data.task.status_label])
        overview.append(["盘点范围部门", data.task.scope_dept_name or "全部部门"])
        overview.append(["盘点范围分类", data.task.scope_category_name or "全部分类"])
        overview.append(["资产总数", data.task.total_assets])
        overview.append(["已盘点", data.task.checked_assets])
        overview.append(["正常", data.task.normal_count])
        overview.append(["异常", data.task.abnormal_count])
        overview.append(["盘亏", data.task.missing_count])
        overview.append(["盘盈", data.task.surplus_count])
        overview.append(["处理单号", data.adjustment.adjustment_no if data.adjustment else "—"])
        overview.append(
            [
                "处理状态",
                data.adjustment.status_label if data.adjustment else "未提交",
            ]
        )
        overview.append(["生成时间", data.generated_at.strftime("%Y-%m-%d %H:%M:%S")])
        self._style_sheet(overview, [22, 40])

        detail = workbook.create_sheet("盘亏与异常明细")
        detail.append(
            ["资产编号", "资产名称", "分类", "部门", "责任人", "位置", "盘点结果", "备注"]
        )
        for item in [*data.missing_items, *data.abnormal_items]:
            detail.append(
                [
                    item.asset_no,
                    item.asset_name,
                    item.category_name,
                    item.dept_name,
                    item.custodian_name,
                    item.actual_location or item.location,
                    item.result_label,
                    item.remark,
                ]
            )
        self._style_sheet(detail, [22, 26, 16, 16, 14, 20, 12, 30])

        surplus = workbook.create_sheet("盘盈明细")
        surplus.append(
            ["资产编号", "资产名称", "分类", "部门", "序列号", "位置", "登记人", "核销状态"]
        )
        for row in data.surplus_items:
            surplus.append(
                [
                    row.asset_no,
                    row.name,
                    row.category_name,
                    row.dept_name,
                    row.serial_no,
                    row.location,
                    row.created_by_name,
                    "已建档" if row.applied_at else "待处理",
                ]
            )
        self._style_sheet(surplus, [22, 26, 16, 16, 20, 20, 14, 12])

        buffer = BytesIO()
        workbook.save(buffer)
        filename = f"盘点报告_{data.task.task_no}.xlsx"
        return InventoryReportExport(
            content=buffer.getvalue(),
            filename=filename,
            media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            row_count=(
                len(data.missing_items)
                + len(data.abnormal_items)
                + len(data.surplus_items)
            ),
        )

    async def list_tasks(
        self,
        context: AuthContext,
        *,
        keyword: str | None,
        status: str | None,
        page: int,
        page_size: int,
    ) -> InventoryTaskPageOut:
        allowed = await self._allowed_depts(context)
        conditions: list = []
        if allowed is not None:
            conditions.append(
                InventoryTask.scope_dept_id.in_(allowed) if allowed else false()
            )
        if keyword:
            like = f"%{keyword.strip()}%"
            conditions.append(
                or_(
                    InventoryTask.task_no.like(like),
                    InventoryTask.name.like(like),
                )
            )
        if status:
            conditions.append(InventoryTask.status == status)
        total = int(
            await self.session.scalar(
                select(func.count()).select_from(InventoryTask).where(*conditions)
            )
            or 0
        )
        rows = list(
            (
                await self.session.execute(
                    select(InventoryTask)
                    .where(*conditions)
                    .order_by(InventoryTask.created_at.desc(), InventoryTask.id.desc())
                    .offset((page - 1) * page_size)
                    .limit(page_size)
                )
            )
            .scalars()
            .all()
        )
        task_ids = [row.id for row in rows]
        stats = await self._stats(task_ids)
        adjustments = await self._latest_adjustments(task_ids)
        return InventoryTaskPageOut(
            items=[
                self._task_out(
                    row,
                    context,
                    allowed,
                    stats.get(row.id),
                    adjustments.get(row.id),
                )
                for row in rows
            ],
            total=total,
            page=page,
            page_size=page_size,
        )

    async def create_task(
        self, payload: InventoryCreate, context: AuthContext
    ) -> InventoryTaskOut:
        allowed = await self._allowed_depts(context)
        dept_ids = await self._resolve_dept_ids(
            payload.scope_dept_id, payload.include_child, allowed
        )
        category_ids = await self._resolve_category_ids(
            payload.scope_category_id, payload.include_child_category
        )

        dept_name = None
        if payload.scope_dept_id is not None:
            dept = await self.session.get(Dept, payload.scope_dept_id)
            dept_name = dept.name if dept else None
        category_name = None
        if payload.scope_category_id is not None:
            category = await self.session.get(AssetCategory, payload.scope_category_id)
            category_name = category.name if category else None

        conditions: list = [
            Asset.deleted_at.is_(None),
            Asset.status.notin_((AssetStatus.SCRAPPED, AssetStatus.PENDING_STOCK)),
        ]
        if dept_ids is not None:
            conditions.append(Asset.dept_id.in_(dept_ids) if dept_ids else false())
        if category_ids is not None:
            conditions.append(
                Asset.category_id.in_(category_ids) if category_ids else false()
            )
        busy_assets = (
            select(InventoryItem.asset_id)
            .join(InventoryTask, InventoryItem.task_id == InventoryTask.id)
            .where(InventoryTask.status == InventoryStatus.IN_PROGRESS)
        )
        conditions.append(Asset.id.notin_(busy_assets))

        assets = list(
            (
                await self.session.execute(
                    select(Asset).where(*conditions).order_by(Asset.id)
                )
            )
            .scalars()
            .all()
        )
        if not assets:
            raise BizError("盘点范围内没有可盘点的资产")

        task = InventoryTask(
            task_no=await self._next_no(),
            name=payload.name,
            status=InventoryStatus.IN_PROGRESS,
            scope_dept_id=payload.scope_dept_id,
            scope_dept_name=dept_name,
            include_child=payload.include_child,
            scope_category_id=payload.scope_category_id,
            scope_category_name=category_name,
            include_child_category=payload.include_child_category,
            planned_date=payload.planned_date,
            remark=payload.remark,
            started_at=now_cst(),
            created_by_id=context.user.id,
            created_by_name=context.user.real_name or context.user.username,
        )
        self.session.add(task)
        await self.session.flush()

        items = [
            InventoryItem(
                task_id=task.id,
                asset_id=asset.id,
                asset_no=asset.asset_no,
                asset_name=asset.name,
                category_name=asset.category.name if asset.category else None,
                dept_name=asset.dept.name if asset.dept else None,
                location=asset.location,
                custodian_name=asset.custodian.real_name if asset.custodian else None,
                expected_status=asset.status,
            )
            for asset in assets
        ]
        self.session.add_all(items)
        await self.audit.record(
            action="asset.inventory.create",
            module="资产盘点",
            summary=f"创建盘点任务：{task.name}（{task.task_no}），共 {len(items)} 项资产",
            actor=context.user,
            resource_type="inventory_task",
            resource_id=task.id,
            after={
                "status": task.status,
                "scope_dept": dept_name,
                "scope_category": category_name,
                "total_assets": len(items),
            },
        )
        await commit_or_conflict(self.session, message=STATUS_CONFLICT_MESSAGE, code=40911)
        return self._task_out(
            task,
            context,
            await self._allowed_depts(context),
            {
                "total": len(items),
                "checked": 0,
                "normal": 0,
                "abnormal": 0,
                "missing": 0,
                "pending": len(items),
            },
        )

    async def get_task(self, task_id: int, context: AuthContext) -> InventoryTaskOut:
        task, allowed = await self._get_task_for_view(task_id, context)
        stats = await self._stats([task.id])
        adjustments = await self._latest_adjustments([task.id])
        return self._task_out(
            task, context, allowed, stats.get(task.id), adjustments.get(task.id)
        )

    async def list_items(
        self,
        task_id: int,
        context: AuthContext,
        *,
        keyword: str | None,
        result: str | None,
        page: int,
        page_size: int,
    ) -> InventoryItemPageOut:
        await self._get_task_for_view(task_id, context)
        conditions: list = [InventoryItem.task_id == task_id]
        if keyword:
            like = f"%{keyword.strip()}%"
            conditions.append(
                or_(
                    InventoryItem.asset_no.like(like),
                    InventoryItem.asset_name.like(like),
                    InventoryItem.custodian_name.like(like),
                )
            )
        if result:
            conditions.append(InventoryItem.result == result)
        total = int(
            await self.session.scalar(
                select(func.count()).select_from(InventoryItem).where(*conditions)
            )
            or 0
        )
        rows = list(
            (
                await self.session.execute(
                    select(InventoryItem)
                    .where(*conditions)
                    .order_by(InventoryItem.id)
                    .offset((page - 1) * page_size)
                    .limit(page_size)
                )
            )
            .scalars()
            .all()
        )
        return InventoryItemPageOut(
            items=[self._item_out(row) for row in rows],
            total=total,
            page=page,
            page_size=page_size,
        )

    async def _get_item(
        self, task_id: int, item_id: int, *, for_update: bool = False
    ) -> InventoryItem:
        stmt = select(InventoryItem).where(
            InventoryItem.id == item_id, InventoryItem.task_id == task_id
        )
        if for_update:
            stmt = stmt.with_for_update()
        item = (await self.session.execute(stmt)).scalar_one_or_none()
        if item is None:
            raise NotFoundError("盘点明细不存在")
        return item

    async def update_item(
        self,
        task_id: int,
        item_id: int,
        payload: InventoryItemUpdate,
        context: AuthContext,
    ) -> InventoryItemOut:
        task, _ = await self._get_task_for_manage(task_id, context)
        item = await self._get_item(task_id, item_id, for_update=True)
        previous = item.result
        item.result = payload.result
        item.actual_location = payload.actual_location
        item.remark = payload.remark
        item.checked_by_id = context.user.id
        item.checked_by_name = context.user.real_name or context.user.username
        item.checked_at = now_cst()
        await self.audit.record(
            action="asset.inventory.item.update",
            module="资产盘点",
            summary=f"盘点核对：{item.asset_no} 标记为"
            f"{INVENTORY_RESULT_LABELS.get(item.result, item.result)}",
            actor=context.user,
            resource_type="inventory_task",
            resource_id=task.id,
            before={"asset_no": item.asset_no, "result": previous},
            after={
                "asset_no": item.asset_no,
                "result": item.result,
                "actual_location": item.actual_location,
                "remark": item.remark,
            },
        )
        await commit_or_conflict(self.session, message=STATUS_CONFLICT_MESSAGE, code=40911)
        return self._item_out(item)

    async def scan(
        self,
        task_id: int,
        payload: InventoryScanPayload,
        context: AuthContext,
    ) -> InventoryItemOut:
        task, _ = await self._get_task_for_manage(task_id, context)
        code = payload.code.strip()
        item = (
            await self.session.execute(
                select(InventoryItem)
                .where(
                    InventoryItem.task_id == task.id,
                    InventoryItem.asset_no == code,
                )
                .with_for_update()
            )
        ).scalar_one_or_none()
        if item is None:
            raise NotFoundError("盘点范围内未找到该资产")
        previous = item.result
        item.result = payload.result
        item.actual_location = payload.actual_location
        item.remark = payload.remark
        item.checked_by_id = context.user.id
        item.checked_by_name = context.user.real_name or context.user.username
        item.checked_at = now_cst()
        await self.audit.record(
            action="asset.inventory.scan",
            module="资产盘点",
            summary=f"扫码盘点：{item.asset_no} 标记为"
            f"{INVENTORY_RESULT_LABELS.get(item.result, item.result)}",
            actor=context.user,
            resource_type="inventory_task",
            resource_id=task.id,
            before={"asset_no": item.asset_no, "result": previous},
            after={"asset_no": item.asset_no, "result": item.result},
        )
        await commit_or_conflict(self.session, message=STATUS_CONFLICT_MESSAGE, code=40911)
        return self._item_out(item)

    async def complete(
        self, task_id: int, context: AuthContext
    ) -> InventoryTaskOut:
        task, allowed = await self._get_task_for_manage(task_id, context)
        stats = (await self._stats([task.id])).get(task.id, {})
        pending = int(stats.get("pending", 0))
        if pending:
            raise BizError(f"还有 {pending} 项资产未盘点，无法完成")
        task.status = InventoryStatus.COMPLETED
        task.completed_at = now_cst()
        await self.audit.record(
            action="asset.inventory.complete",
            module="资产盘点",
            summary=f"完成盘点任务：{task.name}（{task.task_no}）",
            actor=context.user,
            resource_type="inventory_task",
            resource_id=task.id,
            before={"status": InventoryStatus.IN_PROGRESS},
            after={"status": task.status, **stats},
        )
        await commit_or_conflict(self.session, message=STATUS_CONFLICT_MESSAGE, code=40911)
        new_stats = (await self._stats([task.id])).get(task.id)
        return self._task_out(task, context, allowed, new_stats)

    async def cancel(
        self,
        task_id: int,
        payload: InventoryCancelPayload,
        context: AuthContext,
    ) -> InventoryTaskOut:
        task, allowed = await self._get_task_for_manage(task_id, context)
        task.status = InventoryStatus.CANCELLED
        task.cancelled_at = now_cst()
        task.cancel_reason = payload.reason
        await self.audit.record(
            action="asset.inventory.cancel",
            module="资产盘点",
            summary=f"取消盘点任务：{task.name}（{task.task_no}）",
            actor=context.user,
            resource_type="inventory_task",
            resource_id=task.id,
            before={"status": InventoryStatus.IN_PROGRESS},
            after={"status": task.status, "reason": payload.reason},
        )
        await commit_or_conflict(self.session, message=STATUS_CONFLICT_MESSAGE, code=40911)
        return self._task_out(
            task, context, allowed, (await self._stats([task.id])).get(task.id)
        )


__all__ = ["InventoryService"]
