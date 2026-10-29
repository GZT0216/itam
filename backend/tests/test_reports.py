from datetime import timedelta
from io import BytesIO

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import func, insert, select

from app.core.assets import AssetStatus
from app.core.database import AsyncSessionLocal
from app.core.security import hash_password
from app.models import (
    Asset,
    AssetCategory,
    AssetHistory,
    AuditLog,
    Dept,
    Role,
    User,
    user_roles,
)
from app.models.base import now_cst
from tests.helpers import auth_headers, login

TEST_PASSWORD = "Test@12345"


async def _dept_id(code: str) -> int:
    async with AsyncSessionLocal() as session:
        dept = (
            await session.execute(select(Dept).where(Dept.code == code))
        ).scalar_one()
        return dept.id


async def _category_id(code: str) -> int:
    async with AsyncSessionLocal() as session:
        category = (
            await session.execute(
                select(AssetCategory).where(AssetCategory.code == code)
            )
        ).scalar_one()
        return category.id


async def _ensure_user(
    username: str, real_name: str, role_code: str, dept_code: str
) -> int:
    async with AsyncSessionLocal() as session:
        role = (
            await session.execute(select(Role).where(Role.code == role_code))
        ).scalar_one()
        dept = (
            await session.execute(select(Dept).where(Dept.code == dept_code))
        ).scalar_one()
        user = User(
            username=username,
            password_hash=hash_password(TEST_PASSWORD),
            real_name=real_name,
            dept_id=dept.id,
            status=1,
            must_change_password=False,
        )
        session.add(user)
        await session.flush()
        await session.execute(
            insert(user_roles).values(user_id=user.id, role_id=role.id)
        )
        await session.commit()
        return user.id


async def _create_asset(
    asset_no: str,
    name: str,
    *,
    dept_id: int,
    status: str,
    custodian_id: int | None = None,
    location: str | None = "A 座 101",
    purchase_date=None,
    purchase_price=None,
    warranty_end=None,
    serial_no: str | None = None,
    supplier: str | None = "测试供应商",
) -> int:
    async with AsyncSessionLocal() as session:
        category = (
            await session.execute(
                select(AssetCategory).where(AssetCategory.code == "HW-LAPTOP")
            )
        ).scalar_one()
        asset = Asset(
            asset_no=asset_no,
            name=name,
            category_id=category.id,
            dept_id=dept_id,
            custodian_id=custodian_id,
            status=status,
            location=location,
            purchase_date=purchase_date,
            purchase_price=purchase_price,
            warranty_end=warranty_end,
            serial_no=serial_no,
            supplier=supplier,
        )
        session.add(asset)
        await session.commit()
        return asset.id


async def _login_as(client: AsyncClient, username: str) -> dict[str, str]:
    response = await client.post(
        "/api/v1/auth/login",
        json={"username": username, "password": TEST_PASSWORD},
    )
    assert response.status_code == 200, response.text
    return auth_headers(response.json()["data"]["access_token"])


async def test_asset_reports_metrics_idle_compliance_and_export(client: AsyncClient):
    token = (await login(client))["access_token"]
    headers = auth_headers(token)
    today = now_cst().date()
    it_dept = await _dept_id("IT")
    fin_dept = await _dept_id("FIN")

    async with AsyncSessionLocal() as session:
        admin = (
            await session.execute(select(User).where(User.username == "admin"))
        ).scalar_one()
        admin_id = admin.id

    in_use_id = await _create_asset(
        "RPT-ASSET-001",
        "分析用笔记本",
        dept_id=it_dept,
        status=AssetStatus.IN_USE,
        custodian_id=admin_id,
        purchase_date=today - timedelta(days=40),
        purchase_price=5000,
        warranty_end=today + timedelta(days=10),
        serial_no="SN-RPT-001",
    )
    idle_id = await _create_asset(
        "RPT-ASSET-002",
        "闲置笔记本",
        dept_id=fin_dept,
        status=AssetStatus.IDLE,
        location=None,
        purchase_date=None,
        purchase_price=None,
        warranty_end=today - timedelta(days=10),
        serial_no=None,
    )
    await _create_asset(
        "RPT-ASSET-003",
        "待入库笔记本",
        dept_id=it_dept,
        status=AssetStatus.PENDING_STOCK,
        purchase_date=today - timedelta(days=5),
        purchase_price=3000,
        serial_no="SN-RPT-003",
        location="B 座 202",
    )
    await _create_asset(
        "RPT-ASSET-004",
        "已报废笔记本",
        dept_id=fin_dept,
        status=AssetStatus.SCRAPPED,
        location=None,
        purchase_date=None,
        purchase_price=None,
    )

    async with AsyncSessionLocal() as session:
        session.add(
            AssetHistory(
                asset_id=idle_id,
                action="STATUS_CHANGE",
                action_label="状态变更",
                from_status=AssetStatus.IN_USE,
                to_status=AssetStatus.IDLE,
                before_snapshot={},
                after_snapshot={},
                operator_name="系统管理员",
                summary="测试闲置时长",
                created_at=now_cst() - timedelta(days=40),
            )
        )
        await session.commit()

    overview = await client.get("/api/v1/reports/assets/overview", headers=headers)
    assert overview.status_code == 200, overview.text
    data = overview.json()["data"]
    metrics = data["metrics"]
    assert metrics["total_count"] == 4
    assert metrics["total_value"] == 8000.0
    assert metrics["in_use_count"] == 1
    assert metrics["idle_count"] == 1
    assert metrics["warranty_expiring_30d"] == 1
    assert metrics["warranty_expired"] == 1
    assert metrics["compliance_asset_count"] == 1
    status_map = {item["status"]: item for item in data["status_distribution"]}
    assert status_map[AssetStatus.IDLE]["count"] == 1
    dept_names = {item["name"]: item for item in data["dept_distribution"]}
    assert dept_names["信息技术部"]["count"] == 2
    assert len(data["purchase_trend"]) == 12

    assets = await client.get(
        "/api/v1/reports/assets",
        params={"status": AssetStatus.IDLE, "page": 1, "page_size": 10},
        headers=headers,
    )
    assert assets.status_code == 200, assets.text
    assert assets.json()["data"]["total"] == 1
    assert assets.json()["data"]["items"][0]["asset_no"] == "RPT-ASSET-002"

    idle = await client.get(
        "/api/v1/reports/assets/idle",
        params={"page": 1, "page_size": 10},
        headers=headers,
    )
    assert idle.status_code == 200, idle.text
    idle_data = idle.json()["data"]
    assert idle_data["total"] == 1
    assert idle_data["items"][0]["idle_days"] >= 39

    summary = await client.get(
        "/api/v1/reports/assets/compliance/summary", headers=headers
    )
    assert summary.status_code == 200, summary.text
    summary_data = summary.json()["data"]
    counts = {item["issue_type"]: item["count"] for item in summary_data["items"]}
    assert summary_data["asset_count"] == 1
    assert counts["MISSING_LOCATION"] == 1
    assert counts["MISSING_PRICE"] == 1
    assert counts["MISSING_PURCHASE_DATE"] == 1
    assert counts["MISSING_SERIAL"] == 1
    assert counts["EXPIRED_WARRANTY"] == 1
    assert counts["MISSING_CUSTODIAN"] == 0

    compliance = await client.get(
        "/api/v1/reports/assets/compliance",
        params={"issue_type": "MISSING_SERIAL", "page": 1, "page_size": 10},
        headers=headers,
    )
    assert compliance.status_code == 200, compliance.text
    compliance_data = compliance.json()["data"]
    assert compliance_data["total"] == 1
    row = compliance_data["items"][0]
    assert row["asset_no"] == "RPT-ASSET-002"
    assert [item["issue_type"] for item in row["issues"]] == ["MISSING_SERIAL"]

    xlsx = await client.get(
        "/api/v1/reports/assets/export",
        params={"report": "OVERVIEW", "format": "XLSX"},
        headers=headers,
    )
    assert xlsx.status_code == 200, xlsx.text
    assert xlsx.content[:2] == b"PK"
    workbook = load_workbook(BytesIO(xlsx.content), read_only=True)
    assert "概览指标" in workbook.sheetnames
    assert "状态分布" in workbook.sheetnames

    pdf = await client.get(
        "/api/v1/reports/assets/export",
        params={"report": "COMPLIANCE", "format": "PDF"},
        headers=headers,
    )
    assert pdf.status_code == 200, pdf.text
    assert pdf.content[:5] == b"%PDF-"
    assert pdf.headers["x-export-rows"] == "1"

    invalid = await client.get(
        "/api/v1/reports/assets/overview",
        params={"purchase_from": "2026-10-10", "purchase_to": "2026-10-01"},
        headers=headers,
    )
    assert invalid.status_code == 400

    async with AsyncSessionLocal() as session:
        audit_count = await session.scalar(
            select(func.count(AuditLog.id)).where(
                AuditLog.action == "asset.report.export"
            )
        )
        assert audit_count == 2


async def test_report_scope_options_and_permissions(client: AsyncClient):
    await _ensure_user("report_it_admin", "报表 IT 管理员", "DEPT_ADMIN", "IT")
    await _ensure_user("report_employee", "报表普通员工", "EMPLOYEE", "IT")
    it_dept = await _dept_id("IT")
    fin_dept = await _dept_id("FIN")
    await _create_asset(
        "RPT-SCOPE-IT",
        "IT 资产",
        dept_id=it_dept,
        status=AssetStatus.IN_STOCK,
        purchase_date=now_cst().date(),
        purchase_price=1000,
        serial_no="SN-SCOPE-IT",
    )
    await _create_asset(
        "RPT-SCOPE-FIN",
        "财务资产",
        dept_id=fin_dept,
        status=AssetStatus.IN_STOCK,
        purchase_date=now_cst().date(),
        purchase_price=2000,
        serial_no="SN-SCOPE-FIN",
    )

    headers = await _login_as(client, "report_it_admin")
    overview = await client.get("/api/v1/reports/assets/overview", headers=headers)
    assert overview.status_code == 200, overview.text
    assert overview.json()["data"]["metrics"]["total_count"] == 1

    options = await client.get("/api/v1/reports/assets/options", headers=headers)
    assert options.status_code == 200, options.text
    dept_codes = {item["name"] for item in options.json()["data"]["departments"]}
    assert "信息技术部" in dept_codes
    assert "财务部" not in dept_codes

    scope_export = await client.get(
        "/api/v1/reports/assets/export",
        params={"report": "ASSETS", "format": "XLSX"},
        headers=headers,
    )
    assert scope_export.status_code == 200, scope_export.text
    workbook = load_workbook(BytesIO(scope_export.content), read_only=True)
    sheet = workbook["资产清单"]
    values = [str(cell.value or "") for row in sheet.iter_rows() for cell in row]
    assert "RPT-SCOPE-IT" in values
    assert "RPT-SCOPE-FIN" not in values

    employee_headers = await _login_as(client, "report_employee")
    denied = await client.get(
        "/api/v1/reports/assets/overview", headers=employee_headers
    )
    assert denied.status_code == 403
    denied_export = await client.get(
        "/api/v1/reports/assets/export",
        params={"report": "OVERVIEW", "format": "XLSX"},
        headers=employee_headers,
    )
    assert denied_export.status_code == 403
